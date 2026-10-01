"""Quotation templates: how a quote looks on paper, per guarantee-letter issuer.

An admin sets one up in two ways:

- "form": designed here. The letterhead, the issuer's wording, the notes and
  who signs are fields (DESIGN); the print page lays them out.
- "excel": their own Excel form, uploaded as it is. The CRM finds its parts
  by their labels, the way a person would (find()): PATIENT NAME:, CONTACT NO.:,
  ADDRESS:, HOSPITAL/CLINIC, Date Issued, Control No., the products table by
  its headers, GRAND TOTAL, and the "Printed Name/Signature" line. Where a
  label isn't enough, a cell can say exactly what goes in it: {{customer}},
  {{line.qty}}... (PLACEHOLDERS). fill() writes a quote into a copy of it.

Both take a logo and an e-signature (an image uploaded, or drawn on screen).

A quote gets a control number on a template the first time it's printed or
downloaded with it (PREFIX-YYYY-MM-00001), and keeps it.
"""

import io
import re
from copy import copy
from datetime import date, datetime
from typing import Any

from sqlalchemy import select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.config import Settings
from app.extra import fields_for
from app.models import Company, Contact, Issuer, Product, Quote, QuoteControl, QuoteTemplate, StoredFile, Team, User
from app.pricing import sales_settings

MAX_IMAGE_BYTES = 2 * 1024 * 1024
MAX_EXCEL_BYTES = 5 * 1024 * 1024

DESIGN_DEFAULTS: dict[str, Any] = {
    "companyName": "",
    "address": "",
    "phone": "",
    "tin": "",
    "vat": "vat",  # vat | non_vat
    "payableTo": "",
    "heading": "PRICE QUOTATION",
    "customerLabel": "Patient name",
    "certification": (
        "This is to certify that we are in the position to quote and provide the medicines based on the prescription "
        "submitted, at the prices indicated, valid up to thirty (30) days from the date of issuance. "
        "Prices may change without prior notice."
    ),
    "acceptance": "We are willing to accept a guarantee letter from {issuer}.",
    "notes": "Prices indicated are already discounted and at a special price for {issuer}. Other discounts no longer apply.\nOur quotation is free of charge.",
    "signatoryName": "",
    "signatoryTitle": "",
    "signatoryContact": "",
}

PLACEHOLDERS = {
    "control_no": "Control number",
    "quote_no": "Quote number",
    "date": "Date issued",
    "valid_until": "Valid until",
    "customer": "Patient or customer name",
    "phone": "Contact number",
    "address": "Address",
    "hospital": "Hospital or clinic",
    "issuer": "Who pays (the guarantee letter's issuer)",
    "letter_no": "Guarantee letter number",
    "letter_amount": "Guarantee letter amount",
    "team": "Team",
    "prepared_by": "Sales representative",
    "total": "Grand total",
    "signatory": "Authorized representative",
    "logo": "The logo (an image placed at this cell)",
    "signature": "The e-signature (an image placed at this cell)",
    "extra.KEY": "A field added to quotes, by its key",
}
LINE_PLACEHOLDERS = {
    "line.no": "Line number",
    "line.generic": "Generic name",
    "line.brand": "Brand name",
    "line.product": "Product name",
    "line.strength": "Strength",
    "line.form": "Preparation",
    "line.unit": "Unit",
    "line.qty": "Quantity",
    "line.price": "Unit price",
    "line.discount": "Discount per unit",
    "line.net": "Price after discount",
    "line.total": "Line total",
}


def normal(text: Any) -> str:
    return re.sub(r"[^a-z0-9]", "", str(text).lower())


# --- a quote, as a form shows it --------------------------------------------------------------


def issuer_name(db: Session, payer: str | None) -> str:
    if not payer or payer == "self":
        return ""
    found = db.get(Issuer, payer)
    return found.name if found else payer


def extra_by(db: Session, quote: Quote, *words: str) -> str | None:
    """A value from the fields added to quotes, by a word in its key or label."""
    values = quote.extra or {}
    for f in fields_for(db, "quote"):
        name = normal(f"{f.key} {f.label}")
        if any(w in name for w in words) and values.get(f.key) not in (None, ""):
            return str(values[f.key])
    return None


def control_number(db: Session, template: QuoteTemplate, quote: Quote, tries: int = 5) -> str:
    """The quote's control number on this template, given the first time it's asked for.

    Several PCs (each Bishnoi One.exe runs its own API) can print at the same moment.
    The template's row is locked while its next number is taken, so the second
    one waits and gets the number after; a clash that still gets through (a
    database without row locks) is caught by the unique number and retried."""
    found = db.scalar(select(QuoteControl).where(QuoteControl.template_id == template.id, QuoteControl.quote_id == quote.id))
    if found:
        return found.number
    locked = db.scalar(select(QuoteTemplate).where(QuoteTemplate.id == template.id).with_for_update()) or template
    on = (quote.created_at or datetime.now()).strftime("%Y-%m")
    prefix = (locked.control_prefix or sales_settings(db).quote_prefix or "Q").strip("-")
    number = f"{prefix}-{on}-{locked.control_next:0{locked.control_digits}d}"
    # a number another PC already gave out (it hadn't moved the counter yet) is skipped
    while db.scalar(select(QuoteControl.id).where(QuoteControl.template_id == locked.id, QuoteControl.number == number)):
        locked.control_next += 1
        number = f"{prefix}-{on}-{locked.control_next:0{locked.control_digits}d}"
    locked.control_next += 1
    db.add(QuoteControl(template_id=locked.id, quote_id=quote.id, number=number))
    try:
        db.commit()
    except IntegrityError:  # asked twice at once: the other one's number stands, or take the next
        db.rollback()
        if tries <= 1:
            raise
        return control_number(db, template, quote, tries - 1)
    return number


def lines_of(db: Session, quote: Quote) -> list[dict]:
    out = []
    for i, line in enumerate(quote.lines, start=1):
        product = db.get(Product, line.get("productId")) if line.get("productId") else None
        qty = line.get("quantity") or 0
        off = (line.get("promoDiscountCentavos") or 0) + (line.get("discountCentavos") or 0) + (line.get("volumeDiscountCentavos") or 0)
        price = (line.get("unitPriceCentavos") or 0) / 100
        each_off = round(off / qty / 100, 2) if qty else 0
        out.append({
            "no": i,
            "generic": (product.generic_name if product else None) or "",
            "brand": (product.name if product else None) or line.get("productName") or "",
            "product": line.get("productName") or "",
            "strength": (product.strength if product else None) or "",
            "form": (product.form if product else None) or line.get("unit") or "",
            "unit": line.get("unit") or "",
            "qty": qty,
            "free": line.get("freeQuantity") or 0,
            "price": price,
            "discount": each_off or None,
            "net": round(price - each_off, 2),
            "total": (line.get("totalCentavos") or 0) / 100,
        })
    return out


def document(db: Session, quote: Quote, template: QuoteTemplate, settings: Settings) -> dict:
    """Everything a template shows for a quote."""
    contact = db.get(Contact, quote.contact_id) if quote.contact_id else None
    company = db.get(Company, quote.company_id) if quote.company_id else (contact.company if contact else None)
    owner = db.get(User, quote.owner_id) if quote.owner_id else None
    team = db.get(Team, quote.team_id) if quote.team_id else None
    letter = quote.letter or {}
    s = sales_settings(db)
    design = {**DESIGN_DEFAULTS, **(template.design or {})}
    issuer = issuer_name(db, quote.payer) or issuer_name(db, template.payer)
    for key in ("acceptance", "notes", "certification", "heading"):
        design[key] = design[key].replace("{issuer}", issuer or "the issuer")
    design["companyName"] = design["companyName"] or settings.org_name
    design["address"] = design["address"] or (s.company_address or "")
    design["phone"] = design["phone"] or (s.company_contact or "")
    design["tin"] = design["tin"] or (s.company_tin or "")
    design["payableTo"] = design["payableTo"] or design["companyName"]
    address = extra_by(db, quote, "address") or (", ".join(p for p in (company.address, company.city) if p) if company and not contact else None)
    total = (quote.tax or {}).get("totalDueCentavos") or quote.total_centavos
    return {
        "controlNumber": control_number(db, template, quote),
        "quoteNumber": quote.number,
        "date": (quote.created_at.date() if quote.created_at else date.today()).isoformat(),
        "validUntil": quote.valid_until,
        "customer": quote.customer_name,
        "phone": extra_by(db, quote, "contact", "phone", "mobile") or (contact and (contact.mobile or contact.phone)) or (company.phone if company else None) or "",
        "address": address or "",
        "hospital": extra_by(db, quote, "hospital", "clinic", "center") or (company.name if contact and company else "") or "",
        "issuer": issuer,
        "letterNumber": letter.get("number") or "",
        "letterAmount": (letter.get("amountCentavos") or 0) / 100 if letter else None,
        "team": team.name if team else "",
        "preparedBy": owner.full_name if owner else "",
        "lines": lines_of(db, quote),
        "total": total / 100,
        "tax": quote.tax,
        "notes": quote.notes or "",
        "terms": quote.terms or "",
        "extra": quote.extra or {},
        "design": design,
    }


# --- the Excel form ---------------------------------------------------------------------------

LABELS = {
    "customer": ("patientname", "nameofpatient", "customername", "nameofcustomer", "beneficiary", "nameofbeneficiary", "patient", "customer", "name"),
    "phone": ("contactno", "contactnumber", "mobileno", "mobilenumber", "cellphoneno", "telephoneno", "phone", "phoneno"),
    "address": ("address", "homeaddress", "patientaddress"),
    "hospital": ("hospitalcliniccenter", "hospitalclinic", "hospitalarea", "hospital", "clinic"),
    "date": ("dateissued", "dateofquotation", "quotationdate", "date"),
    "valid_until": ("validuntil", "validity"),
}
FIELD_NAMES = {
    "customer": "Patient or customer name",
    "phone": "Contact number",
    "address": "Address",
    "hospital": "Hospital or clinic",
    "date": "Date issued",
    "valid_until": "Valid until",
    "control": "Control number",
}
COLUMNS = {
    "no": ("no", "item", "itemno"),
    "generic": ("genericname", "generic", "molecule"),
    "brand": ("brandname", "brand"),
    "product": ("product", "productname", "description", "medicine", "particulars", "itemdescription"),
    "strength": ("strength", "dosagestrength", "dosage"),
    "form": ("preparation", "dosageform", "form"),
    "unit": ("unit", "uom", "packsize"),
    "qty": ("quantity", "qty"),
    "price": ("price", "unitprice", "priceperpiece", "priceperunit", "srp"),
    "discount": ("discount", "less"),
    "net": ("priceafterdiscount", "netprice", "discountedprice"),
    "total": ("total", "amount", "totalamount", "linetotal"),
}
COLUMN_NAMES = {
    "no": "No.", "generic": "Generic name", "brand": "Brand name", "product": "Product", "strength": "Strength",
    "form": "Preparation", "unit": "Unit", "qty": "Quantity", "price": "Price", "discount": "Discount",
    "net": "Price after discount", "total": "Total",
}
PLACEHOLDER = re.compile(r"\{\{\s*([a-zA-Z_.0-9]+)\s*\}\}")


def load(data: bytes, **kw):
    from openpyxl import load_workbook

    return load_workbook(io.BytesIO(data), **kw)


def sheet_of(book, name: str | None):
    if name and name in book.sheetnames:
        return book[name]
    for ws in book.worksheets:
        if "quot" in ws.title.lower():
            return ws
    return book.worksheets[0]


def texts(ws) -> list[tuple[int, int, str]]:
    out = []
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.strip():
                out.append((c.row, c.column, c.value))
    return out


def label_of(text: str) -> str | None:
    t = normal(text)
    if not t or len(t) > 30:
        return None
    for field, names in LABELS.items():
        if t in names:
            return field
    return None


def find(ws) -> dict:
    """Where the form's parts are: {"fields": {field: "C12"}, "table": {...}, "signature": {...}, "logo": "A1"}."""
    from openpyxl.utils import get_column_letter

    cells = texts(ws)
    out: dict[str, Any] = {"fields": {}, "table": None, "signature": None, "logo": None, "placeholders": []}

    # placeholders first: they say exactly where things go
    line_row = None
    line_cols: dict[str, str] = {}
    for r, c, v in cells:
        for name in PLACEHOLDER.findall(v):
            out["placeholders"].append(name)
            if name.startswith("line."):
                line_row = r
                line_cols[name[5:]] = get_column_letter(c)
            elif name == "logo":
                out["logo"] = f"{get_column_letter(c)}{r}"
            elif name == "signature":
                out["signature"] = {"anchor": f"{get_column_letter(c)}{r}", "name": None}

    # the products table: a row of known headers, down to the total
    if line_row:
        total = next(((r, c) for r, c, v in cells if r > line_row and "total" in normal(v) and "{{" not in v), None)
        out["table"] = {"header": line_row - 1, "first": line_row, "last": line_row, "columns": line_cols,
                        "totalRow": total[0] if total else None, "totalCell": None, "fromPlaceholders": True}
    else:
        best = None
        for row in sorted({r for r, _, _ in cells}):
            found = {}
            for r, c, v in cells:
                if r != row:
                    continue
                t = normal(v)
                for key, names in COLUMNS.items():
                    if t in names and key not in found:
                        found[key] = get_column_letter(c)
                        break
            if len(found) >= 3 and ("qty" in found or "total" in found) and (best is None or len(found) > len(best[1])):
                best = (row, found)
        if best:
            header, cols = best
            total_row = next(
                (r for r, c, v in sorted(cells) if r > header and ("grandtotal" in normal(v) or normal(v) in ("total", "totalamount", "totalamountdue", "amountdue"))),
                None,
            )
            if total_row:
                last = max(header + 1, total_row - 1)
            else:  # no total: the sample rows are the ones with something in them
                filled = [c.row for row in ws.iter_rows(min_row=header + 1, max_row=header + 50) for c in row if c.value not in (None, "")]
                last = max([header + 1, *filled])
            total_cell = None
            if total_row:
                col = cols.get("total")
                if col:
                    total_cell = f"{col}{total_row}"
            out["table"] = {"header": header, "first": header + 1, "last": last, "columns": cols, "totalRow": total_row,
                            "totalCell": total_cell, "fromPlaceholders": False}

    top = out["table"]["header"] if out["table"] else ws.max_row + 1

    # labelled fields above the table: the value is to the right, on that row or just below
    for r, c, v in cells:
        if r >= top:
            continue
        if normal(v).startswith("controlno"):
            out["fields"].setdefault("control", f"{get_column_letter(c)}{r}")
            continue
        field = label_of(v)
        if not field or field in out["fields"]:
            continue
        spot = None
        for rr in range(r, min(r + 4, top)):
            for cc in range(c + 1, c + 9):
                val = ws.cell(row=rr, column=cc).value
                if val in (None, "") or (isinstance(val, str) and not val.strip()):
                    continue
                if isinstance(val, str) and (label_of(val) or val.strip().endswith(":") or normal(val) in ("areacode", "dateissued")):
                    break  # the next label: nothing for this one on this row
                spot = (rr, cc)
                break
            if spot:
                break
            if rr > r and any(rr2 == rr and cc2 <= c + 1 and label_of(vv) and not normal(vv).startswith("center") for rr2, cc2, vv in cells):
                break  # reached the next label down
        if spot is None:
            spot = (r, c + 1)
        out["fields"][field] = f"{get_column_letter(spot[1])}{spot[0]}"

    # the signature line: a caption with "signature", under the signatory's name
    if out["signature"] is None:
        for r, c, v in cells:
            t = normal(v)
            if "signature" in t and r > top:
                above = [(rr, cc) for rr, cc, vv in cells if cc == c and r - 4 <= rr < r]
                name = max(above) if above else None  # the nearest line above the caption
                name_row = name[0] if name else r - 1
                anchor_row = max(1, name_row - 2)
                out["signature"] = {"anchor": f"{get_column_letter(c)}{anchor_row}", "name": f"{get_column_letter(c)}{name_row}"}
                break
    if out["logo"] is None:
        out["logo"] = "A1"
    return out


def describe(found: dict) -> dict:
    """What find() found, in words for the admin: what was found where, and what wasn't."""
    have, missing = [], []
    for key, label in FIELD_NAMES.items():
        if key in found["fields"]:
            have.append(f"{label}: {found['fields'][key]}")
        elif key in ("customer", "date", "control"):
            missing.append(label)
    t = found.get("table")
    if t:
        cols = ", ".join(COLUMN_NAMES.get(k, k) for k in t["columns"])
        rows = f"row {t['first']}" if t["first"] == t["last"] else f"rows {t['first']} to {t['last']}"
        have.append(f"Products table: {rows} ({cols})")
        if t.get("totalCell"):
            have.append(f"Grand total: {t['totalCell']}")
        elif not t.get("fromPlaceholders"):
            missing.append("Grand total")
    else:
        missing.append("Products table")
    if found.get("signature"):
        have.append(f"E-signature: at {found['signature']['anchor']}")
    have.append(f"Logo: at {found.get('logo') or 'A1'}")
    return {"found": have, "missing": missing}


def _shift_rows(ws, at: int, count: int) -> None:
    """Make room for more product lines: insert rows at `at`, moving what's below
    (merged cells, row heights, pictures) down with them."""
    ws.insert_rows(at, count)
    for mr in list(ws.merged_cells.ranges):
        if mr.min_row >= at:
            mr.shift(0, count)
    heights = {r: d.height for r, d in ws.row_dimensions.items() if r >= at and d.height}
    for r in sorted(heights, reverse=True):
        ws.row_dimensions[r + count].height = heights[r]
        ws.row_dimensions[r].height = None
    for img in ws._images:
        anchor = img.anchor
        start = getattr(anchor, "_from", None)
        if start is not None and start.row >= at - 1:
            start.row += count
            end = getattr(anchor, "to", None)
            if end is not None:
                end.row += count


def _move(ref: str | None, last: int, by: int) -> str | None:
    if not ref or not by:
        return ref
    m = re.fullmatch(r"([A-Z]+)(\d+)", ref)
    if not m:
        return ref
    row = int(m.group(2))
    return f"{m.group(1)}{row + by if row > last else row}"


def _top_left(ws, ref: str):
    """A merged cell is written at its top-left."""
    for mr in ws.merged_cells.ranges:
        if ref in mr:
            return ws.cell(row=mr.min_row, column=mr.min_col)
    return ws[ref]


def _picture(data: bytes, max_w: int, max_h: int):
    from openpyxl.drawing.image import Image
    from PIL import Image as Pil

    pic = Pil.open(io.BytesIO(data))
    pic.load()
    scale = min(max_w / pic.width, max_h / pic.height, 1)
    out = io.BytesIO()
    pic.convert("RGBA").save(out, "PNG")
    out.seek(0)
    img = Image(out)
    img.width, img.height = int(pic.width * scale), int(pic.height * scale)
    return img


def fill(data: bytes, sheet: str | None, doc: dict, logo: bytes | None, signature: bytes | None) -> bytes:
    """The Excel form, with the quote written in. Only the form's sheet is kept."""
    book = load(data)
    ws = sheet_of(book, sheet)
    for other in [s for s in book.worksheets if s is not ws]:
        book.remove(other)
    found = find(ws)
    lines = doc["lines"]
    table = found.get("table")
    moved = 0
    if table:
        first, last = table["first"], table["last"]
        room = last - first + 1
        if len(lines) > room:
            moved = len(lines) - room
            _shift_rows(ws, last + 1, moved)
            for r in range(last + 1, last + 1 + moved):
                for col in range(1, ws.max_column + 1):
                    src, dst = ws.cell(row=first, column=col), ws.cell(row=r, column=col)
                    if src.has_style:
                        dst._style = copy(src._style)
                if ws.row_dimensions[first].height:
                    ws.row_dimensions[r].height = ws.row_dimensions[first].height
        from openpyxl.cell.cell import MergedCell
        from openpyxl.utils import column_index_from_string

        cols = {k: column_index_from_string(v) for k, v in table["columns"].items()}
        span = range(min(cols.values()), max(cols.values()) + 1) if cols else range(0)
        for i in range(max(room, len(lines))):
            r = first + i
            for col in span:
                cell = ws.cell(row=r, column=col)
                if not isinstance(cell, MergedCell):
                    cell.value = None
            if i >= len(lines):
                continue
            line = lines[i]
            for key, col in cols.items():
                value = line.get(key)
                if key == "product" and "brand" not in cols and "generic" not in cols:
                    value = " ".join(p for p in (line["generic"], line["brand"], line["strength"]) if p) or line["product"]
                ws.cell(row=r, column=col).value = value if value not in ("", None) else None
        if table.get("totalCell"):
            _top_left(ws, _move(table["totalCell"], last, moved)).value = doc["total"]
        last_row = last
    else:
        last_row = ws.max_row

    values = {
        "customer": doc["customer"], "phone": doc["phone"], "address": doc["address"], "hospital": doc["hospital"],
        "date": date.fromisoformat(doc["date"]), "valid_until": date.fromisoformat(doc["validUntil"]) if doc["validUntil"] else None,
    }
    for field, ref in found["fields"].items():
        cell = _top_left(ws, _move(ref, last_row, moved))
        if field == "control":
            text = cell.value if isinstance(cell.value, str) else "Control No."
            m = re.match(r"(\s*control\s*no\.?\s*:?\s*)", text, re.I)
            cell.value = f"{m.group(1) if m else 'Control No. '}{doc['controlNumber']}"
        elif field in values:
            cell.value = values[field]

    # placeholders anywhere
    words = {
        "control_no": doc["controlNumber"], "quote_no": doc["quoteNumber"], "date": values["date"], "valid_until": values["valid_until"],
        "customer": doc["customer"], "phone": doc["phone"], "address": doc["address"], "hospital": doc["hospital"],
        "issuer": doc["issuer"], "letter_no": doc["letterNumber"], "letter_amount": doc["letterAmount"], "team": doc["team"],
        "prepared_by": doc["preparedBy"], "total": doc["total"], "signatory": doc["design"].get("signatoryName") or "",
        "logo": None, "signature": None,
    }
    for key, value in (doc.get("extra") or {}).items():
        words[f"extra.{key}"] = value
    for row in ws.iter_rows():
        for cell in row:
            if not isinstance(cell.value, str) or "{{" not in cell.value:
                continue
            names = PLACEHOLDER.findall(cell.value)
            if len(names) == 1 and PLACEHOLDER.fullmatch(cell.value.strip()) and not names[0].startswith("line."):
                cell.value = words.get(names[0])
            else:
                cell.value = PLACEHOLDER.sub(lambda m: "" if words.get(m.group(1)) is None else str(words.get(m.group(1))), cell.value) or None

    if logo:
        ws.add_image(_picture(logo, 220, 70), _move(found.get("logo") or "A1", last_row, moved))
    if signature and found.get("signature"):
        ws.add_image(_picture(signature, 200, 60), _move(found["signature"]["anchor"], last_row, moved))
        name = found["signature"].get("name")
        signatory = doc["design"].get("signatoryName")
        if name and signatory:
            _top_left(ws, _move(name, last_row, moved)).value = signatory

    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


# --- stored files -----------------------------------------------------------------------------

IMAGE_KINDS = {b"\x89PNG": ("image/png", "png"), b"\xff\xd8\xff": ("image/jpeg", "jpg")}


def image_kind(data: bytes) -> tuple[str, str] | None:
    return next((v for magic, v in IMAGE_KINDS.items() if data.startswith(magic)), None)


def stored(db: Session, ref: str | None) -> StoredFile | None:
    return db.get(StoredFile, ref) if ref else None
