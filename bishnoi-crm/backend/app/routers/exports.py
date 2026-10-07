"""Exports: customers, with everything on their form, as an Excel workbook
(a second sheet lists their contacts) or a CSV. The columns and how values
are written are app/customer_sheet.py's, the same an import reads back.

Only customers within the person's reach, with the customer list's filters.
Exporting needs customers.view and reports.export: a whole customer list is
the company's to give out, so reps don't hold it by default. Every export is
in the audit log. A blank sheet (template=true, headers only) needs only
customers.create: it's for filling in and importing.
"""

import csv
import io
from datetime import date
from typing import Literal

from fastapi import APIRouter, Query, Response
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from sqlalchemy import func, or_, select

from app.access import CurrentAccess
from app.audit import audit
from app.customer_sheet import CONTACT_HEADERS, Lookups, added_fields, contact_rows, headers, row_of
from app.deps import DB, CurrentUser
from app.models import Company
from app.scope import within

router = APIRouter(prefix="/exports", tags=["exports"])

MAX_ROWS = 20000


def workbook(sheets: list[tuple[str, list[str], list[list[str]]]]) -> bytes:
    book = Workbook(write_only=True)
    for title, head, rows in sheets:
        sheet = book.create_sheet(title)
        sheet.freeze_panes = "A2"
        for i, h in enumerate(head):
            sheet.column_dimensions[get_column_letter(i + 1)].width = max(12, min(40, len(h) + 4))
        bold = Font(bold=True)
        cells = []
        for h in head:
            c = WriteOnlyCell(sheet, value=h)
            c.font = bold
            cells.append(c)
        sheet.append(cells)
        for row in rows:
            sheet.append(row)
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def as_csv(head: list[str], rows: list[list[str]]) -> bytes:
    out = io.StringIO()
    writer = csv.writer(out)
    writer.writerow(head)
    writer.writerows(rows)
    return ("﻿" + out.getvalue()).encode("utf-8")  # the mark makes Excel read it as UTF-8


XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@router.get("/customers")
def export_customers(
    access: CurrentAccess,
    me: CurrentUser,
    db: DB,
    file_format: Literal["xlsx", "csv"] = Query("xlsx", alias="format"),
    template: bool = False,
    q: str | None = Query(None, max_length=100),
    status_: str | None = Query(None, alias="status", max_length=40),
    owner: str | None = Query(None, max_length=36),
) -> Response:
    head = headers(db)
    stamp = date.today().isoformat()
    if template:
        access.require("customers.create")
        # The import's columns only: Team and Added on aren't read.
        head = [h for h in head if h not in ("Team", "Added on")]
        body = workbook([("Customers", head, [])]) if file_format == "xlsx" else as_csv(head, [])
        return file_out(body, file_format, "customers-blank")

    access.require("customers.view")
    access.require("reports.export")
    query = within(select(Company), Company, access)
    if q:
        like = f"%{q.lower()}%"
        query = query.where(or_(func.lower(Company.name).like(like), func.lower(Company.code).like(like), func.lower(Company.city).like(like), func.lower(Company.industry).like(like)))
    if status_:
        query = query.where(Company.status == status_)
    if owner:
        query = query.where(Company.owner_id == (access.user_id if owner == "me" else owner))
    companies = db.scalars(query.order_by(func.lower(Company.name)).limit(MAX_ROWS)).unique().all()
    fields, look = added_fields(db), Lookups(db)
    rows = [row_of(c, fields, look) for c in companies]
    if file_format == "xlsx":
        contacts = contact_rows(db, access, {c.id for c in companies}, unlinked=not (q or status_ or owner))
        body = workbook([("Customers", head, rows), ("Contacts", CONTACT_HEADERS, contacts)])
    else:
        body = as_csv(head, rows)
    audit(db, me, "customers.exported", target=f"{len(rows)} customers ({file_format})", target_type="company")
    db.commit()
    return file_out(body, file_format, f"customers-{stamp}")


def file_out(body: bytes, file_format: str, name: str) -> Response:
    return Response(
        body,
        media_type=XLSX if file_format == "xlsx" else "text/csv; charset=utf-8",
        headers={"Content-Disposition": f'attachment; filename="{name}.{file_format}"', "Cache-Control": "private, no-store"},
    )
