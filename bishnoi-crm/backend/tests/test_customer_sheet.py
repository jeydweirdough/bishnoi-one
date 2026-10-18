"""Customers out to a spreadsheet and back in: every field on their form,
written as people read it, read back the same way, through the form's checks."""

import io

from openpyxl import load_workbook

from app.models import FormField
from tests.test_customers import company, contact, crew  # noqa: F401 (crew is a fixture)
from tests.test_imports import read, xlsx


def fields(app) -> None:
    """A pick-list with short codes and full names, and a link to another customer."""
    with app.state.sessions() as db:
        db.add_all([
            FormField(entity="company", key="x_dept", label="Department", type="choice", options=["MED ONCO", "HEMA"], option_names={"MED ONCO": "Medical oncology", "HEMA": "Hematology"}, sort_order=1),
            FormField(entity="company", key="x_hosp", label="Hospital name", type="customer", sort_order=2),
        ])
        db.commit()


def sheet(data: bytes, name: str = "Customers") -> list[list]:
    book = load_workbook(io.BytesIO(data), read_only=True)
    return [[c if c is not None else "" for c in row] for row in book[name].iter_rows(values_only=True)]


def test_customers_export_with_their_fields_readable(app, client, crew, ana):
    fields(app)
    hospital = company(client, ana, "St. Luke's")
    doctor = company(client, ana, "Dr. Cruz", extra={"x_dept": "MED ONCO", "x_hosp": hospital["id"]})
    contact(client, ana, "Maria", companyId=hospital["id"], email="maria@stlukes.test")
    contact(client, ana, "Lone")  # with no customer: in a full export, not a filtered one

    r = client.get("/exports/customers", headers=ana)
    assert r.status_code == 200, r.text
    assert "attachment" in r.headers["content-disposition"] and r.headers["content-disposition"].endswith('.xlsx"')
    rows = sheet(r.content)
    head = rows[0]
    assert head[:3] == ["Customer code", "Name", "Type"] and head[-2:] == ["Department", "Hospital name"]
    cruz = next(row for row in rows if row[1] == "Dr. Cruz")
    assert cruz[0] == doctor["code"] and cruz[-2:] == ["MED ONCO", "St. Luke's"]  # the hospital by name, not its id
    assert cruz[head.index("Assigned to")] == "ana@example.com"
    people = sheet(r.content, "Contacts")
    assert [row[:3] for row in people[1:]] == [["", "", "Lone"], [hospital["code"], "St. Luke's", "Maria"]]
    assert [row[2] for row in sheet(client.get("/exports/customers?q=st", headers=ana).content, "Contacts")[1:]] == ["Maria"]

    csv_out = client.get("/exports/customers?format=csv&q=cruz", headers=ana)
    assert csv_out.status_code == 200 and csv_out.text.count("\n") == 2  # the header and Dr. Cruz: the list's filter applies

    # A whole customer list is the company's to give out: reps can't export it, but can get a blank sheet.
    assert client.get("/exports/customers", headers=crew["ria"]).status_code == 403
    blank = client.get("/exports/customers?template=true", headers=crew["ria"])
    assert blank.status_code == 200 and sheet(blank.content) == [[h for h in head if h not in ("Team", "Added on")]]


def test_customers_import_adds_updates_and_says_why_a_row_is_skipped(app, client, crew, ana):
    fields(app)
    hospital = company(client, ana, "St. Luke's")
    doctor = company(client, ana, "Dr. Cruz", extra={"x_dept": "MED ONCO"})
    rows = [
        ["Code", "Customer Name", "Department", "Hospital name", "Assigned to"],
        [doctor["code"], "Dr. Cruz", "Hematology", "St. Luke's", ""],  # an update: the full name of a pick
        ["", "Dr. Reyes", "MED ONCO", hospital["code"], "ria@example.com"],  # new, given to Ria
        ["CRM-999999", "Nobody", "", "", ""],  # no such code
        ["", "Dr. Lim", "Cardiology", "", ""],  # not an option
        ["", "Dr. Tan", "", "", "nobody@example.com"],  # nobody here
    ]
    got = read(client, ana, "customers", "c.xlsx", xlsx(rows)).json()
    m = got["mapping"]
    assert (m["code"], m["name"], m["extra:x_dept"], m["extra:x_hosp"], m["owner"]) == (0, 1, 2, 3, 4)

    body = {"mapping": m, "rows": got["rows"], "dryRun": True}
    check = client.post("/imports/customers", json=body, headers=ana).json()
    assert (check["added"], check["updated"]) == (1, 1), check
    assert [p["row"] for p in check["problems"]] == [4, 5, 6]
    assert "CRM-999999" in check["problems"][0]["message"] and "Cardiology" in check["problems"][1]["message"]
    assert client.get("/companies?q=reyes", headers=ana).json()["total"] == 0  # a check saves nothing

    done = client.post("/imports/customers", json={**body, "dryRun": False}, headers=ana).json()
    assert (done["added"], done["updated"], len(done["problems"])) == (1, 1, 3)
    cruz = client.get(f"/companies/{doctor['id']}", headers=ana).json()
    assert cruz["extra"]["x_dept"] == "HEMA" and cruz["extra"]["x_hosp"] == hospital["id"]
    reyes = client.get("/companies?q=reyes", headers=ana).json()["items"][0]
    assert reyes["owner"]["name"] == "Ria" and reyes["extra"] == {"x_dept": "MED ONCO", "x_hosp": hospital["id"]}
    assert reyes["code"]  # given a code like any new customer

    # A rep updates only what's theirs: Dr. Cruz is Ana's.
    mine = client.post("/imports/customers", json={**body, "rows": [got["rows"][0]], "dryRun": True}, headers=crew["ria"]).json()
    assert mine["updated"] == 0 and "can change" in mine["problems"][0]["message"]


def test_an_exported_file_imports_back_unchanged(app, client, ana):
    fields(app)
    hospital = company(client, ana, "St. Luke's")
    company(client, ana, "Dr. Cruz", extra={"x_dept": "HEMA", "x_hosp": hospital["id"]}, city="Quezon City")
    out = client.get("/exports/customers", headers=ana).content
    got = read(client, ana, "customers", "back.xlsx", out).json()
    done = client.post("/imports/customers", json={"mapping": got["mapping"], "rows": got["rows"], "dryRun": False}, headers=ana).json()
    assert (done["added"], done["updated"], done["problems"]) == (0, 2, [])
    assert sheet(client.get("/exports/customers", headers=ana).content) == sheet(out)
