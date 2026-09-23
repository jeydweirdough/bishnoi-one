"""Reading a spreadsheet somebody actually has.

Nobody keeps their products in the shape Orbit would like. So Orbit does not
ask for a template: it reads whatever workbook or CSV is handed to it, says
which columns it found, guesses which is which, and lets a person correct the
guesses before anything is written.

That guess is the whole trick. "Item Code", "item_code", "CODE" and "SKU" all
mean the same thing to a person, so they should mean the same thing here.
"""

import csv
import io
import re
from dataclasses import dataclass, field
from typing import Any

from openpyxl import load_workbook

MAX_ROWS = 20_000
PREVIEW_ROWS = 5


class SheetError(ValueError):
    """Worded for the person who picked the file."""


@dataclass
class Sheet:
    columns: list[str]
    rows: list[dict[str, str]]
    # Sheets a workbook holds beyond the one that was read, so a person who
    # picked the wrong tab can be told what else is in there.
    other_tabs: list[str] = field(default_factory=list)


def tidy(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def unique_headers(raw: list[str]) -> list[str]:
    """Column names as the file has them, with the blanks and the repeats made
    usable so a mapping can point at exactly one column."""
    seen: dict[str, int] = {}
    out = []
    for i, name in enumerate(raw):
        header = tidy(name) or f"Column {i + 1}"
        if header in seen:
            seen[header] += 1
            header = f"{header} ({seen[header]})"
        else:
            seen[header] = 1
        out.append(header)
    return out


def with_rows(rows: list[dict[str, str]]) -> list[dict[str, str]]:
    """A file with headings and nothing under them is not an import, and saying
    so beats accepting it and reporting that nothing happened."""
    if not rows:
        raise SheetError("That file has column headings but no rows under them.")
    return rows


def read_csv(data: bytes) -> Sheet:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        # Excel on a Philippine or European machine saves CSV in the local code
        # page, which is not UTF-8 and will not decode as it.
        text = data.decode("cp1252", errors="replace")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text), dialect)
    rows = [r for r in reader if any(tidy(c) for c in r)]
    if not rows:
        raise SheetError("That file has nothing in it.")
    headers = unique_headers(rows[0])
    body = [dict(zip(headers, [tidy(c) for c in r], strict=False)) for r in rows[1 : MAX_ROWS + 1]]
    return Sheet(headers, with_rows(body))


def read_excel(data: bytes) -> Sheet:
    try:
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises several different things for a bad file
        raise SheetError("That doesn't open as a spreadsheet. Save it as .xlsx or .csv and try again.") from exc
    try:
        page = book[book.sheetnames[0]]
        grid = [[tidy(c) for c in row] for row in page.iter_rows(max_row=MAX_ROWS + 1, values_only=True)]
        others = book.sheetnames[1:]
    finally:
        book.close()

    grid = [r for r in grid if any(r)]
    if not grid:
        raise SheetError("That sheet has nothing in it.")
    headers = unique_headers(grid[0])
    return Sheet(headers, with_rows([dict(zip(headers, r, strict=False)) for r in grid[1:]]), others)


def read_sheet(filename: str, data: bytes) -> Sheet:
    if filename.lower().endswith((".csv", ".txt")):
        return read_csv(data)
    if filename.lower().endswith((".xlsx", ".xlsm")):
        return read_excel(data)
    if filename.lower().endswith(".xls"):
        raise SheetError("That's the old Excel format. Open it and save as .xlsx, then try again.")
    raise SheetError("Pick a .xlsx or .csv file.")


def normalise(header: str) -> str:
    return re.sub(r"[^a-z0-9]+", "", header.casefold())


def guess(columns: list[str], wanted: dict[str, tuple[str, ...]]) -> dict[str, str | None]:
    """Which column probably holds which field: an exact match on the tidied
    name first, then a column that contains the word. A field with no obvious
    column is left for the person to point at."""
    taken: set[str] = set()
    found: dict[str, str | None] = {}
    by_shape = {normalise(c): c for c in columns}

    for field_key, aliases in wanted.items():
        hit = next((by_shape[normalise(a)] for a in aliases if normalise(a) in by_shape and by_shape[normalise(a)] not in taken), None)
        if hit is None:
            hit = next(
                (c for c in columns if c not in taken and any(normalise(a) and normalise(a) in normalise(c) for a in aliases)),
                None,
            )
        if hit is not None:
            taken.add(hit)
        found[field_key] = hit
    return found
