"""Remove all residual test products from Bishnoi Inventory and Bishnoi CRM.
Keeps ONLY the 100 products listed in Pricelist02.2026.
Targets:
1. Inventory PostgreSQL: neondb_getmeds, neondb
2. Inventory SQLite: inventory-getmeds.db, inventory.db
3. CRM PostgreSQL: neondb_getmeds, neondb
4. CRM SQLite: crm-getmeds.db, crm.db
"""

import io
import os
import re
import shutil
import sqlite3
import sys
import openpyxl
import psycopg

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

EXCEL_PATH = r"C:\Users\Getmeds Backend\Downloads\Price List 2026 Finance (1).xlsx"

INV_GETMEDS_URL = "postgresql://neondb_owner:npg_5TPJAcRI4XZt@ep-floral-bonus-b3j059si-pooler.c-4.ap-southeast-1.aws.neon.tech/neondb_getmeds?sslmode=require"
INV_BASE_URL = "postgresql://neondb_owner:npg_5TPJAcRI4XZt@ep-floral-bonus-b3j059si-pooler.c-4.ap-southeast-1.aws.neon.tech/neondb?sslmode=require"

CRM_GETMEDS_URL = "postgresql://neondb_owner:npg_K4FmOfRkalv9@ep-jolly-sun-b3r8zkb6-pooler.c-4.ap-southeast-1.aws.neon.tech/neondb_getmeds?sslmode=require"
CRM_BASE_URL = "postgresql://neondb_owner:npg_K4FmOfRkalv9@ep-jolly-sun-b3r8zkb6-pooler.c-4.ap-southeast-1.aws.neon.tech/neondb?sslmode=require"

INV_GETMEDS_SQLITE = r"C:\Users\Getmeds Backend\Desktop\bishnoi-one\bishnoi-inventory\backend\.data\inventory-getmeds.db"
INV_BASE_SQLITE = r"C:\Users\Getmeds Backend\Desktop\bishnoi-one\bishnoi-inventory\backend\.data\inventory.db"

CRM_GETMEDS_SQLITE = r"C:\Users\Getmeds Backend\Desktop\bishnoi-one\bishnoi-crm\backend\.data\crm-getmeds.db"
CRM_BASE_SQLITE = r"C:\Users\Getmeds Backend\Desktop\bishnoi-one\bishnoi-crm\backend\.data\crm.db"

def make_backup(file_path):
    if not os.path.exists(file_path):
        return
    backup_dir = os.path.join(os.path.dirname(file_path), "backups")
    os.makedirs(backup_dir, exist_ok=True)
    from datetime import datetime
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    base_name = os.path.basename(file_path)
    backup_path = os.path.join(backup_dir, f"{base_name}_{ts}.bak")
    shutil.copy2(file_path, backup_path)
    print(f"Backup created: {backup_path}")

def clean_sku(brand, generic, strength):
    s = f"{brand} {strength}".strip()
    s = re.sub(r'\[.*?\]|\(.*?\)', '', s)
    s = re.sub(r'[^a-zA-Z0-9]+', '-', s).strip('-').upper()
    parts = s.split('-')
    dedup = []
    for p in parts:
        if not dedup or p != dedup[-1]:
            dedup.append(p)
    res = "-".join(dedup)
    if not res:
        res = re.sub(r'[^a-zA-Z0-9]+', '-', generic).strip('-').upper()
    if len(res) > 35:
        res = res[:35].rstrip('-')
    return res

def get_excel_skus():
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    ws = wb['Pricelist02.2026']
    skus = []
    for r in range(3, 103):
        generic = str(ws.cell(r, 4).value or '').strip()
        brand = str(ws.cell(r, 5).value or '').strip()
        strength = str(ws.cell(r, 6).value or '').strip()
        sku = clean_sku(brand, generic, strength)
        if sku in skus:
            sku = f"{sku[:35]}-{r}"
        skus.append(sku)
    return skus

def clean_inventory_postgres(db_url, db_label, excel_skus):
    print(f"\n--- Cleaning Inventory PostgreSQL ({db_label}) ---")
    with psycopg.connect(db_url) as conn:
        with conn.cursor() as cur:
            cur.execute("SELECT id FROM nbin_products WHERE NOT (sku = ANY(%s))", (excel_skus,))
            res_ids = [r[0] for r in cur.fetchall()]
            print(f"Found {len(res_ids)} residual products in {db_label}")
            if not res_ids:
                cur.execute("SELECT count(*) FROM nbin_products")
                print(f"Total products in {db_label}: {cur.fetchone()[0]}")
                return

            cur.execute("DELETE FROM nbin_movements WHERE product_id = ANY(%s)", (res_ids,))
            cur.execute("DELETE FROM nbin_serials WHERE product_id = ANY(%s)", (res_ids,))
            cur.execute("DELETE FROM nbin_order_lines WHERE product_id = ANY(%s)", (res_ids,))
            cur.execute("DELETE FROM nbin_orders WHERE id NOT IN (SELECT order_id FROM nbin_order_lines)")
            cur.execute("DELETE FROM nbin_purchase_order_lines WHERE product_id = ANY(%s)", (res_ids,))
            cur.execute("DELETE FROM nbin_purchase_orders WHERE id NOT IN (SELECT order_id FROM nbin_purchase_order_lines)")
            cur.execute("DELETE FROM nbin_transfer_lines WHERE product_id = ANY(%s)", (res_ids,))
            cur.execute("DELETE FROM nbin_transfers WHERE id NOT IN (SELECT transfer_id FROM nbin_transfer_lines)")
            cur.execute("DELETE FROM nbin_adjustment_lines WHERE product_id = ANY(%s)", (res_ids,))
            cur.execute("DELETE FROM nbin_adjustments WHERE id NOT IN (SELECT adjustment_id FROM nbin_adjustment_lines)")
            cur.execute("DELETE FROM nbin_receipt_lines WHERE product_id = ANY(%s)", (res_ids,))
            cur.execute("DELETE FROM nbin_receipts WHERE id NOT IN (SELECT receipt_id FROM nbin_receipt_lines)")
            cur.execute("DELETE FROM nbin_batches WHERE product_id = ANY(%s)", (res_ids,))
            cur.execute("DELETE FROM nbin_supplier_products WHERE product_id = ANY(%s)", (res_ids,))
            cur.execute("DELETE FROM nbin_stock WHERE product_id = ANY(%s)", (res_ids,))
            cur.execute("DELETE FROM nbin_product_images WHERE product_id = ANY(%s)", (res_ids,))
            cur.execute("DELETE FROM nbin_products WHERE id = ANY(%s)", (res_ids,))

            conn.commit()

            cur.execute("SELECT count(*) FROM nbin_products")
            print(f"After clean in {db_label}: Products count = {cur.fetchone()[0]}")

def clean_inventory_sqlite(db_path, db_name, excel_skus):
    print(f"\n--- Cleaning Inventory SQLite ({db_name}: {db_path}) ---")
    if not os.path.exists(db_path):
        print(f"Skipping {db_path} (does not exist)")
        return
    make_backup(db_path)

    conn = sqlite3.connect(db_path)
    c = conn.cursor()

    ph = ','.join('?' * len(excel_skus))
    c.execute(f"SELECT id FROM nbin_products WHERE sku NOT IN ({ph})", excel_skus)
    res_ids = [r[0] for r in c.fetchall()]
    print(f"Found {len(res_ids)} residual products in SQLite {db_name}")

    if res_ids:
        ph_ids = ','.join('?' * len(res_ids))
        c.execute(f"DELETE FROM nbin_movements WHERE product_id IN ({ph_ids})", res_ids)
        c.execute(f"DELETE FROM nbin_serials WHERE product_id IN ({ph_ids})", res_ids)
        c.execute(f"DELETE FROM nbin_order_lines WHERE product_id IN ({ph_ids})", res_ids)
        c.execute("DELETE FROM nbin_orders WHERE id NOT IN (SELECT order_id FROM nbin_order_lines)")
        c.execute(f"DELETE FROM nbin_purchase_order_lines WHERE product_id IN ({ph_ids})", res_ids)
        c.execute("DELETE FROM nbin_purchase_orders WHERE id NOT IN (SELECT order_id FROM nbin_purchase_order_lines)")
        c.execute(f"DELETE FROM nbin_transfer_lines WHERE product_id IN ({ph_ids})", res_ids)
        c.execute("DELETE FROM nbin_transfers WHERE id NOT IN (SELECT transfer_id FROM nbin_transfer_lines)")
        c.execute(f"DELETE FROM nbin_adjustment_lines WHERE product_id IN ({ph_ids})", res_ids)
        c.execute("DELETE FROM nbin_adjustments WHERE id NOT IN (SELECT adjustment_id FROM nbin_adjustment_lines)")
        c.execute(f"DELETE FROM nbin_receipt_lines WHERE product_id IN ({ph_ids})", res_ids)
        c.execute("DELETE FROM nbin_receipts WHERE id NOT IN (SELECT receipt_id FROM nbin_receipt_lines)")
        c.execute(f"DELETE FROM nbin_batches WHERE product_id IN ({ph_ids})", res_ids)
        c.execute(f"DELETE FROM nbin_supplier_products WHERE product_id IN ({ph_ids})", res_ids)
        c.execute(f"DELETE FROM nbin_stock WHERE product_id IN ({ph_ids})", res_ids)
        c.execute(f"DELETE FROM nbin_product_images WHERE product_id IN ({ph_ids})", res_ids)
        c.execute(f"DELETE FROM nbin_products WHERE id IN ({ph_ids})", res_ids)
        conn.commit()

    c.execute("SELECT count(*) FROM nbin_products")
    print(f"After clean in SQLite {db_name}: Products count = {c.fetchone()[0]}")
    conn.close()

def clean_crm_postgres(db_url, db_label, excel_skus):
    print(f"\n--- Cleaning CRM PostgreSQL ({db_label}) ---")
    with psycopg.connect(db_url) as conn:
        with conn.cursor() as cur:
            cur.execute("SELECT id FROM nbcm_products WHERE NOT (code = ANY(%s))", (excel_skus,))
            res_ids = [r[0] for r in cur.fetchall()]
            print(f"Found {len(res_ids)} residual products in CRM {db_label}")

            cur.execute("DELETE FROM nbcm_bundles WHERE code IN ('KIDS-KIT', 'MAINT-3')")
            cur.execute("DELETE FROM nbcm_promos WHERE name IN ('Buy 10, get 1', 'Flu season 10% off', 'Volume discount')")

            if res_ids:
                cur.execute("DELETE FROM nbcm_product_prices WHERE product_id = ANY(%s)", (res_ids,))
                cur.execute("DELETE FROM nbcm_team_prices WHERE product_id = ANY(%s)", (res_ids,))
                cur.execute("DELETE FROM nbcm_price_conditions WHERE product_id = ANY(%s)", (res_ids,))
                cur.execute("DELETE FROM nbcm_products WHERE id = ANY(%s)", (res_ids,))

            conn.commit()

            cur.execute("SELECT count(*) FROM nbcm_products")
            print(f"After clean in CRM {db_label}: Products count = {cur.fetchone()[0]}")

def clean_crm_sqlite(db_path, db_name, excel_skus):
    print(f"\n--- Cleaning CRM SQLite ({db_name}: {db_path}) ---")
    if not os.path.exists(db_path):
        print(f"Skipping {db_path} (does not exist)")
        return
    make_backup(db_path)

    conn = sqlite3.connect(db_path)
    c = conn.cursor()

    ph = ','.join('?' * len(excel_skus))
    c.execute(f"SELECT id FROM nbcm_products WHERE code NOT IN ({ph})", excel_skus)
    res_ids = [r[0] for r in c.fetchall()]
    print(f"Found {len(res_ids)} residual products in CRM SQLite {db_name}")

    c.execute("DELETE FROM nbcm_bundles WHERE code IN ('KIDS-KIT', 'MAINT-3')")
    c.execute("DELETE FROM nbcm_promos WHERE name IN ('Buy 10, get 1', 'Flu season 10% off', 'Volume discount')")

    if res_ids:
        ph_ids = ','.join('?' * len(res_ids))
        c.execute(f"DELETE FROM nbcm_product_prices WHERE product_id IN ({ph_ids})", res_ids)
        c.execute(f"DELETE FROM nbcm_team_prices WHERE product_id IN ({ph_ids})", res_ids)
        c.execute(f"DELETE FROM nbcm_price_conditions WHERE product_id IN ({ph_ids})", res_ids)
        c.execute(f"DELETE FROM nbcm_products WHERE id IN ({ph_ids})", res_ids)

    conn.commit()

    c.execute("SELECT count(*) FROM nbcm_products")
    print(f"After clean in CRM SQLite {db_name}: Products count = {c.fetchone()[0]}")
    conn.close()

def main():
    print("=================================================================")
    print("      REMOVE ALL RESIDUAL TEST PRODUCTS FROM INVENTORY & CRM     ")
    print("=================================================================")

    excel_skus = get_excel_skus()
    print(f"Valid Excel Products in Pricelist02.2026: {len(excel_skus)}")

    # 1. Clean Inventory
    clean_inventory_postgres(INV_GETMEDS_URL, "neondb_getmeds", excel_skus)
    clean_inventory_postgres(INV_BASE_URL, "neondb", excel_skus)
    clean_inventory_sqlite(INV_GETMEDS_SQLITE, "inventory-getmeds.db", excel_skus)
    clean_inventory_sqlite(INV_BASE_SQLITE, "inventory.db", excel_skus)

    # 2. Clean CRM
    clean_crm_postgres(CRM_GETMEDS_URL, "neondb_getmeds", excel_skus)
    clean_crm_postgres(CRM_BASE_URL, "neondb", excel_skus)
    clean_crm_sqlite(CRM_GETMEDS_SQLITE, "crm-getmeds.db", excel_skus)
    clean_crm_sqlite(CRM_BASE_SQLITE, "crm.db", excel_skus)

    print("\n=================================================================")
    print("        ALL RESIDUAL PRODUCTS REMOVED SUCCESSFULLY               ")
    print("=================================================================")

if __name__ == '__main__':
    main()
