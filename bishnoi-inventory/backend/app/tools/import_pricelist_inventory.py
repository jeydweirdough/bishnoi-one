"""Import products from Pricelist02.2026 into:
1. Bishnoi Inventory (Postgres & SQLite): 100 products, 0 stock across all warehouses.
2. Bishnoi CRM (Postgres & SQLite): 100 products, 4 price tiers (DOCTOR, PATIENT, SRP, TRADE),
   plus 29 Bundles and 39 B2C Promos.

Source:
- Price List 2026 Finance (1).xlsx (Sheet: "Pricelist02.2026", "SEPT PROMO BUNDLES", "B2C Promo")
"""

import io
import json
import os
import re
import shutil
import sqlite3
import sys
import uuid
from datetime import datetime, UTC
import openpyxl
import psycopg

# Set console output encoding to utf-8
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

EXCEL_PATH = r"C:\Users\Getmeds Backend\Downloads\Price List 2026 Finance (1).xlsx"

# Neon URLs
INV_GETMEDS_URL = "postgresql://neondb_owner:npg_5TPJAcRI4XZt@ep-floral-bonus-b3j059si-pooler.c-4.ap-southeast-1.aws.neon.tech/neondb_getmeds?sslmode=require"
INV_BASE_URL = "postgresql://neondb_owner:npg_5TPJAcRI4XZt@ep-floral-bonus-b3j059si-pooler.c-4.ap-southeast-1.aws.neon.tech/neondb?sslmode=require"

CRM_GETMEDS_URL = "postgresql://neondb_owner:npg_K4FmOfRkalv9@ep-jolly-sun-b3r8zkb6-pooler.c-4.ap-southeast-1.aws.neon.tech/neondb_getmeds?sslmode=require"
CRM_BASE_URL = "postgresql://neondb_owner:npg_K4FmOfRkalv9@ep-jolly-sun-b3r8zkb6-pooler.c-4.ap-southeast-1.aws.neon.tech/neondb?sslmode=require"

# SQLite paths
INV_GETMEDS_SQLITE = r"C:\Users\Getmeds Backend\Desktop\bishnoi-one\bishnoi-inventory\backend\.data\inventory-getmeds.db"
INV_BASE_SQLITE = r"C:\Users\Getmeds Backend\Desktop\bishnoi-one\bishnoi-inventory\backend\.data\inventory.db"
CRM_GETMEDS_SQLITE = r"C:\Users\Getmeds Backend\Desktop\bishnoi-one\bishnoi-crm\backend\.data\crm-getmeds.db"
CRM_BASE_SQLITE = r"C:\Users\Getmeds Backend\Desktop\bishnoi-one\bishnoi-crm\backend\.data\crm.db"

def make_backup(file_path):
    if not os.path.exists(file_path):
        return
    backup_dir = os.path.join(os.path.dirname(file_path), "backups")
    os.makedirs(backup_dir, exist_ok=True)
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    base_name = os.path.basename(file_path)
    backup_path = os.path.join(backup_dir, f"{base_name}_{ts}.bak")
    shutil.copy2(file_path, backup_path)
    print(f"Backup created: {backup_path}")

def clean_sku(brand, generic, strength):
    s = f"{brand} {strength}".strip()
    s = re.sub(r'\[.*?\]|\(.*?\)', '', s)  # Remove brackets/parentheses
    s = re.sub(r'[^a-zA-Z0-9]+', '-', s).strip('-').upper()
    parts = s.split('-')
    dedup = []
    for p in parts:
        if not dedup or p != dedup[-1]:
            dedup.append(p)
    res = "-".join(dedup)
    if not res:
        res = re.sub(r'[^a-zA-Z0-9]+', '-', generic).strip('-').upper()
    if len(res) > 35:
        res = res[:35].rstrip('-')
    return res

def to_centavos(val):
    if val is None:
        return 0
    try:
        if isinstance(val, (int, float)):
            return int(round(float(val) * 100))
        cleaned = re.sub(r'[^0-9.]', '', str(val))
        return int(round(float(cleaned) * 100)) if cleaned else 0
    except:
        return 0

def load_excel_data():
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    ws_pl = wb['Pricelist02.2026']

    skus = set()
    products = []

    unit_map = {
        'ampoule': 'ampule',
        'ampule': 'ampule',
        'vial': 'vial',
        'box': 'box',
        'piece': 'pc',
        'unit': 'pc',
        'pouch': 'pack',
        'pack': 'pack',
        'bottle': 'bottle',
        'tablet': 'tablet',
        'capsule': 'capsule',
        'sachet': 'sachet',
        'set': 'set',
    }

    for r in range(3, 103):
        account = str(ws_pl.cell(r, 1).value or '').strip()
        category = str(ws_pl.cell(r, 2).value or '').strip()
        tax = str(ws_pl.cell(r, 3).value or '').strip()
        generic = str(ws_pl.cell(r, 4).value or '').strip()
        brand = str(ws_pl.cell(r, 5).value or '').strip()
        strength = str(ws_pl.cell(r, 6).value or '').strip()
        form = str(ws_pl.cell(r, 7).value or '').strip()
        unit_raw = str(ws_pl.cell(r, 8).value or '').strip()
        pack_size = str(ws_pl.cell(r, 9).value or '').strip()
        count = ws_pl.cell(r, 11).value
        shelf_life = str(ws_pl.cell(r, 12).value or '').strip()
        doc_u = ws_pl.cell(r, 13).value
        doc_p = ws_pl.cell(r, 14).value
        pat_u = ws_pl.cell(r, 15).value
        pat_p = ws_pl.cell(r, 16).value
        srp_u = ws_pl.cell(r, 17).value
        srp_p = ws_pl.cell(r, 18).value
        dist_u = ws_pl.cell(r, 19).value
        dist_p = ws_pl.cell(r, 20).value
        trade_u = ws_pl.cell(r, 22).value
        trade_p = ws_pl.cell(r, 23).value

        sku = clean_sku(brand, generic, strength)
        if sku in skus:
            sku = f"{sku[:35]}-{r}"
        skus.add(sku)

        prod_uuid = str(uuid.uuid5(uuid.NAMESPACE_DNS, f"bishnoi.product.{sku}"))
        unit_code = unit_map.get(unit_raw.lower(), 'pc')
        is_vat = 'VAT' in tax.upper() and 'NON' not in tax.upper()

        prod_obj = {
            'id': prod_uuid,
            'sku': sku,
            'brand': brand,
            'generic': generic,
            'strength': strength,
            'form': form,
            'unit_raw': unit_raw or 'Piece',
            'unit_code': unit_code,
            'pack_size': pack_size,
            'count': int(count) if count and isinstance(count, (int, float)) else 1,
            'shelf_life': shelf_life,
            'category': category,
            'tax': tax,
            'is_vat': is_vat,
            'account': account,
            'doc_u_cents': to_centavos(doc_u),
            'doc_p_cents': to_centavos(doc_p),
            'pat_u_cents': to_centavos(pat_u),
            'pat_p_cents': to_centavos(pat_p),
            'srp_u_cents': to_centavos(srp_u),
            'srp_p_cents': to_centavos(srp_p),
            'dist_u_cents': to_centavos(dist_u),
            'dist_p_cents': to_centavos(dist_p),
            'trade_u_cents': to_centavos(trade_u),
            'trade_p_cents': to_centavos(trade_p),
        }
        products.append(prod_obj)

    # 2. Bundles
    ws_sb = wb['SEPT PROMO BUNDLES']
    bundles = []
    current_bundle = None

    for r in range(5, 65):
        b_name = ws_sb.cell(r, 1).value
        contents = ws_sb.cell(r, 2).value
        offer_type = ws_sb.cell(r, 3).value
        promo_p = ws_sb.cell(r, 4).value

        if b_name and 'bundle' in str(b_name).lower():
            if current_bundle:
                bundles.append(current_bundle)
            price_cents = to_centavos(promo_p)
            current_bundle = {
                'id': str(uuid.uuid5(uuid.NAMESPACE_DNS, f"bishnoi.bundle.{str(b_name).strip()}")),
                'code': re.sub(r'[^a-zA-Z0-9]+', '-', str(b_name)).strip('-').upper(),
                'name': str(b_name).strip(),
                'offer_type': str(offer_type or '').strip(),
                'price_cents': price_cents,
                'items_raw': []
            }

        if contents and current_bundle:
            item_str = str(contents).strip().rstrip('+').strip()
            if item_str:
                current_bundle['items_raw'].append(item_str)

    if current_bundle:
        bundles.append(current_bundle)

    # Resolve items
    for b in bundles:
        parsed_items = []
        for raw_item in b['items_raw']:
            m = re.match(r'^(\d+)(?:\+(\d+))?\s+(?:FREE\s+)?(.*)$', raw_item, re.IGNORECASE)
            qty = 1
            search_name = raw_item
            if m:
                qty = int(m.group(1))
                if m.group(2):
                    qty += int(m.group(2))
                search_name = m.group(3).strip()

            matched_prod = None
            clean_search = re.sub(r'[^a-zA-Z0-9]', '', search_name.lower())
            for p in products:
                brand_clean = re.sub(r'[^a-zA-Z0-9]', '', p['brand'].lower())
                gen_clean = re.sub(r'[^a-zA-Z0-9]', '', p['generic'].lower())
                if brand_clean in clean_search or clean_search in brand_clean:
                    matched_prod = p
                    break
                if gen_clean and (gen_clean in clean_search or clean_search in gen_clean):
                    matched_prod = p
                    break

            if matched_prod:
                parsed_items.append({
                    'productId': matched_prod['id'],
                    'sku': matched_prod['sku'],
                    'name': matched_prod['brand'],
                    'qty': qty
                })
            else:
                parsed_items.append({
                    'productId': None,
                    'name': search_name,
                    'qty': qty
                })
        b['items'] = parsed_items

    # 3. B2C Promos
    ws_b2c = wb['B2C Promo']
    promos = []
    for r in range(5, 44):
        brand_b2c = str(ws_b2c.cell(r, 1).value or '').strip()
        promo_p = ws_b2c.cell(r, 6).value

        if not brand_b2c:
            continue

        matched_id = None
        for p in products:
            if p['brand'].lower() == brand_b2c.lower() or brand_b2c.lower() in p['brand'].lower():
                matched_id = p['id']
                break

        promos.append({
            'id': str(uuid.uuid5(uuid.NAMESPACE_DNS, f"bishnoi.promo.b2c.{brand_b2c}")),
            'name': f"B2C 20% Off - {brand_b2c}",
            'kind': 'percent',
            'percent_bp': 2000,
            'product_id': matched_id,
            'brand': brand_b2c,
            'promo_price_cents': to_centavos(promo_p),
        })

    return products, bundles, promos

def sync_inventory_postgres(db_url, db_label, products):
    print(f"\n--- Synchronizing Bishnoi Inventory PostgreSQL ({db_label}) ---")
    now = datetime.now(UTC)

    with psycopg.connect(db_url) as conn:
        with conn.cursor() as cur:
            # 1. Warehouses & Locations
            cur.execute("SELECT id, name FROM nbin_warehouses")
            warehouses = cur.fetchall()
            if not warehouses:
                wh_id = str(uuid.uuid4())
                cur.execute("INSERT INTO nbin_warehouses (id, code, name, is_active, created_at, updated_at) VALUES (%s, 'MAIN', 'Main Warehouse', true, %s, %s)", (wh_id, now, now))
                warehouses = [(wh_id, 'Main Warehouse')]

            cur.execute("SELECT id, warehouse_id FROM nbin_locations")
            loc_rows = cur.fetchall()
            loc_by_wh = {}
            for loc_id, wh_id in loc_rows:
                if wh_id not in loc_by_wh:
                    loc_by_wh[wh_id] = loc_id

            for wh_id, wh_name in warehouses:
                if wh_id not in loc_by_wh:
                    loc_id = str(uuid.uuid4())
                    cur.execute("INSERT INTO nbin_locations (id, warehouse_id, code, name, is_default, is_active) VALUES (%s, %s, 'MAIN', 'Main Location', true, true)", (loc_id, wh_id))
                    loc_by_wh[wh_id] = loc_id

            # 2. Categories
            categories_needed = {p['category'] for p in products if p['category']}
            cur.execute("SELECT name, id FROM nbin_categories")
            existing_cats = dict(cur.fetchall())

            for cat_name in categories_needed:
                if cat_name not in existing_cats:
                    cat_id = str(uuid.uuid5(uuid.NAMESPACE_DNS, f"bishnoi.category.{cat_name}"))
                    cur.execute("""
                        INSERT INTO nbin_categories (id, name, is_active, created_at, updated_at)
                        VALUES (%s, %s, true, %s, %s)
                    """, (cat_id, cat_name, now, now))
                    existing_cats[cat_name] = cat_id

            # 3. Brands
            brands_needed = {'2MG', 'Getmeds'}
            cur.execute("SELECT name, id FROM nbin_brands")
            existing_brands = dict(cur.fetchall())
            for b_name in brands_needed:
                if b_name not in existing_brands:
                    b_id = str(uuid.uuid5(uuid.NAMESPACE_DNS, f"bishnoi.brand.{b_name}"))
                    cur.execute("""
                        INSERT INTO nbin_brands (id, name, is_active, created_at, updated_at)
                        VALUES (%s, %s, true, %s, %s)
                    """, (b_id, b_name, now, now))
                    existing_brands[b_name] = b_id

            # 4. Units
            cur.execute("SELECT code FROM nbin_units")
            existing_units = {r[0] for r in cur.fetchall()}
            for p in products:
                uc = p['unit_code']
                if uc not in existing_units:
                    cur.execute("""
                        INSERT INTO nbin_units (code, name, is_active)
                        VALUES (%s, %s, true)
                    """, (uc, p['unit_raw'].title()))
                    existing_units.add(uc)

            # 5. Products
            cur.execute("SELECT sku, id FROM nbin_products")
            existing_prods = dict(cur.fetchall())

            prod_created = 0
            prod_updated = 0

            for p in products:
                p_id = p['id']
                sku = p['sku']
                name = p['brand']
                desc = f"{p['generic']} {p['strength']} {p['form']}".strip()
                cat_id = existing_cats.get(p['category'])
                brand_id = existing_brands.get(p['account'], existing_brands.get('Getmeds'))
                p_type = 'stock'
                u_code = p['unit_code']
                cost_p = p['dist_u_cents'] or p['dist_p_cents']
                sell_p = p['srp_u_cents'] or p['pat_u_cents'] or p['srp_p_cents']
                tax_bp = 1200 if p['is_vat'] else 0

                attrs = json.dumps({
                    'account': p['account'],
                    'category': p['category'],
                    'generic_name': p['generic'],
                    'dosage_strength': p['strength'],
                    'dosage_form': p['form'],
                    'pack_size': p['pack_size'],
                    'count': p['count'],
                    'shelf_life': p['shelf_life'],
                    'doctor_price_unit': p['doc_u_cents'],
                    'doctor_price_pack': p['doc_p_cents'],
                    'patient_price_unit': p['pat_u_cents'],
                    'patient_price_pack': p['pat_p_cents'],
                    'srp_unit': p['srp_u_cents'],
                    'srp_pack': p['srp_p_cents'],
                    'distributor_price_unit': p['dist_u_cents'],
                    'distributor_price_pack': p['dist_p_cents'],
                    'trade_price_unit': p['trade_u_cents'],
                    'trade_price_pack': p['trade_p_cents'],
                })

                if sku in existing_prods:
                    db_id = existing_prods[sku]
                    cur.execute("""
                        UPDATE nbin_products
                        SET name = %s, description = %s, category_id = %s, brand_id = %s, product_type = %s,
                            unit_code = %s, cost_price = %s, selling_price = %s, tax_rate_bp = %s,
                            attributes = %s, status = 'active', updated_at = %s
                        WHERE id = %s
                    """, (name, desc, cat_id, brand_id, p_type, u_code, cost_p, sell_p, tax_bp, attrs, now, db_id))
                    prod_updated += 1
                else:
                    cur.execute("""
                        INSERT INTO nbin_products (
                            id, sku, name, description, category_id, brand_id, product_type,
                            unit_code, cost_price, selling_price, tax_rate_bp, tax_inclusive,
                            min_stock, reorder_point, reorder_qty, status, track_batches,
                            track_expiry, track_serials, attributes, created_at, updated_at
                        ) VALUES (
                            %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, true,
                            0, 0, 0, 'active', true, true, false, %s, %s, %s
                        )
                    """, (p_id, sku, name, desc, cat_id, brand_id, p_type, u_code, cost_p, sell_p, tax_bp, attrs, now, now))
                    prod_created += 1
                    existing_prods[sku] = p_id

            # 6. Set 0 Stock for all products across warehouses
            # Instruction: "Products first, put them into 0 stocks first."
            cur.execute("SELECT product_id, warehouse_id, location_id, id FROM nbin_stock")
            existing_stocks = {(r[0], r[1], r[2]): r[3] for r in cur.fetchall()}

            stock_created = 0
            for p in products:
                prod_id = existing_prods.get(p['sku'], p['id'])
                for wh_id, _ in warehouses:
                    loc_id = loc_by_wh[wh_id]
                    key = (prod_id, wh_id, loc_id)
                    if key not in existing_stocks:
                        stock_id = str(uuid.uuid4())
                        cur.execute("""
                            INSERT INTO nbin_stock (id, product_id, warehouse_id, location_id, on_hand, damaged, updated_at)
                            VALUES (%s, %s, %s, %s, 0, 0, %s)
                        """, (stock_id, prod_id, wh_id, loc_id, now))
                        stock_created += 1
                        existing_stocks[key] = stock_id

            conn.commit()
            print(f"Inventory {db_label}: {prod_created} products created, {prod_updated} updated. {stock_created} zero-stock records added (Total products: {len(existing_prods)}).")

def sync_crm_postgres(db_url, db_label, products, bundles, promos):
    print(f"\n--- Synchronizing Bishnoi CRM PostgreSQL ({db_label}) ---")
    now = datetime.now(UTC)

    with psycopg.connect(db_url) as conn:
        with conn.cursor() as cur:
            # 1. Price Types
            cur.execute("SELECT code FROM nbcm_price_types")
            existing_pts = {r[0] for r in cur.fetchall()}

            needed_pts = [
                ('DOCTOR', "Doctor's Price", 1),
                ('PATIENT', "Patient Price", 2),
                ('SRP', "Suggested Retail Price (SRP)", 3),
                ('TRADE', "Drugstore / Hospital Trade Price", 4),
            ]
            for pt_code, pt_name, pt_sort in needed_pts:
                if pt_code not in existing_pts:
                    cur.execute("""
                        INSERT INTO nbcm_price_types (code, name, entry, rule, requires_notes, needs_approval, is_active, sort_order)
                        VALUES (%s, %s, 'manual', 'custom', false, false, true, %s)
                    """, (pt_code, pt_name, pt_sort))
                    existing_pts.add(pt_code)

            # 2. Products
            cur.execute("SELECT code, id FROM nbcm_products")
            existing_crm_prods = dict(cur.fetchall())

            crm_created = 0
            crm_updated = 0
            prod_code_to_id = {}

            for p in products:
                sku = p['sku']
                name = p['brand']
                gen = p['generic']
                st = p['strength']
                form = p['form']
                unit = p['unit_raw']
                vat_cat = 'vat' if p['is_vat'] else 'exempt'
                inv_id = p['id']

                extra_json = json.dumps({
                    'account': p['account'],
                    'category': p['category'],
                    'pack_size': p['pack_size'],
                    'count': p['count'],
                    'shelf_life': p['shelf_life'],
                    'distributor_price_unit': p['dist_u_cents'],
                    'distributor_price_pack': p['dist_p_cents'],
                })

                if sku in existing_crm_prods:
                    u_id = existing_crm_prods[sku]
                    cur.execute("""
                        UPDATE nbcm_products
                        SET name = %s, generic_name = %s, strength = %s, form = %s, unit = %s,
                            is_active = true, vat_category = %s, inventory_id = %s, extra = %s, updated_at = %s
                        WHERE id = %s
                    """, (name, gen, st, form, unit, vat_cat, inv_id, extra_json, now, u_id))
                    crm_updated += 1
                    prod_code_to_id[sku] = u_id
                else:
                    u_id = str(uuid.uuid4())
                    cur.execute("""
                        INSERT INTO nbcm_products (id, code, name, generic_name, strength, form, unit, is_active, vat_category, inventory_id, extra, created_at, updated_at)
                        VALUES (%s, %s, %s, %s, %s, %s, %s, true, %s, %s, %s, %s, %s)
                    """, (u_id, sku, name, gen, st, form, unit, vat_cat, inv_id, extra_json, now, now))
                    crm_created += 1
                    existing_crm_prods[sku] = u_id
                    prod_code_to_id[sku] = u_id

                # Upsert Prices
                prices_to_set = [
                    ('DOCTOR', p['doc_u_cents'] or p['doc_p_cents']),
                    ('PATIENT', p['pat_u_cents'] or p['pat_p_cents']),
                    ('SRP', p['srp_u_cents'] or p['srp_p_cents']),
                    ('TRADE', p['trade_u_cents'] or p['trade_p_cents']),
                ]
                for pt_code, amt in prices_to_set:
                    if amt > 0:
                        cur.execute("""
                            INSERT INTO nbcm_product_prices (product_id, price_type_code, amount_centavos)
                            VALUES (%s, %s, %s)
                            ON CONFLICT (product_id, price_type_code)
                            DO UPDATE SET amount_centavos = EXCLUDED.amount_centavos
                        """, (u_id, pt_code, amt))

            # 3. Bundles
            cur.execute("SELECT code, id FROM nbcm_bundles")
            existing_bundles = dict(cur.fetchall())
            bundle_created = 0
            bundle_updated = 0

            for b in bundles:
                b_code = b['code']
                b_name = b['name']
                b_price = b['price_cents']

                mapped_items = []
                for it in b['items']:
                    it_sku = it.get('sku')
                    p_id = prod_code_to_id.get(it_sku) if it_sku else None
                    mapped_items.append({
                        'productId': p_id,
                        'name': it['name'],
                        'qty': it['qty']
                    })

                items_json = json.dumps(mapped_items)

                if b_code in existing_bundles:
                    b_id = existing_bundles[b_code]
                    cur.execute("""
                        UPDATE nbcm_bundles
                        SET name = %s, price_centavos = %s, items = %s, is_active = true
                        WHERE id = %s
                    """, (b_name, b_price, items_json, b_id))
                    bundle_updated += 1
                else:
                    b_id = str(uuid.uuid4())
                    cur.execute("""
                        INSERT INTO nbcm_bundles (id, code, name, price_centavos, as_one_line, items, is_active, created_at, team_ids)
                        VALUES (%s, %s, %s, %s, false, %s, true, %s, '[]'::json)
                    """, (b_id, b_code, b_name, b_price, items_json, now))
                    bundle_created += 1
                    existing_bundles[b_code] = b_id

            # 4. Promos
            cur.execute("SELECT name, id FROM nbcm_promos")
            existing_promos = dict(cur.fetchall())
            promo_created = 0
            promo_updated = 0

            for pr in promos:
                pr_name = pr['name']
                pr_kind = pr['kind']
                pr_pct = pr['percent_bp']
                target_p_id = pr.get('product_id')
                p_ids = [target_p_id] if target_p_id else []
                p_ids_json = json.dumps(p_ids)

                if pr_name in existing_promos:
                    pr_id = existing_promos[pr_name]
                    cur.execute("""
                        UPDATE nbcm_promos
                        SET kind = %s, percent_bp = %s, product_ids = %s, is_active = true
                        WHERE id = %s
                    """, (pr_kind, pr_pct, p_ids_json, pr_id))
                    promo_updated += 1
                else:
                    pr_id = str(uuid.uuid4())
                    cur.execute("""
                        INSERT INTO nbcm_promos (id, name, kind, percent_bp, tiers, product_ids, is_active, created_at, team_ids)
                        VALUES (%s, %s, %s, %s, '[]'::json, %s, true, %s, '[]'::json)
                    """, (pr_id, pr_name, pr_kind, pr_pct, p_ids_json, now))
                    promo_created += 1
                    existing_promos[pr_name] = pr_id

            conn.commit()
            print(f"CRM {db_label}: {crm_created} products created, {crm_updated} updated. {bundle_created} bundles created, {bundle_updated} updated. {promo_created} promos created, {promo_updated} updated.")

def sync_inventory_sqlite(db_path, db_name, products):
    print(f"\n--- Synchronizing Inventory SQLite ({db_name}: {db_path}) ---")
    if not os.path.exists(db_path):
        print(f"Skipping {db_path} (does not exist)")
        return
    make_backup(db_path)
    now = datetime.now(UTC).strftime("%Y-%m-%d %H:%M:%S")

    conn = sqlite3.connect(db_path)
    c = conn.cursor()

    # Warehouses
    c.execute("SELECT id, name FROM nbin_warehouses")
    warehouses = c.fetchall()
    if not warehouses:
        wh_id = str(uuid.uuid4())
        c.execute("INSERT INTO nbin_warehouses (id, code, name, is_active, created_at, updated_at) VALUES (?, 'MAIN', 'Main Warehouse', 1, ?, ?)", (wh_id, now, now))
        warehouses = [(wh_id, 'Main Warehouse')]

    c.execute("SELECT id, warehouse_id FROM nbin_locations")
    loc_by_wh = {}
    for loc_id, wh_id in c.fetchall():
        if wh_id not in loc_by_wh:
            loc_by_wh[wh_id] = loc_id

    for wh_id, _ in warehouses:
        if wh_id not in loc_by_wh:
            loc_id = str(uuid.uuid4())
            c.execute("INSERT INTO nbin_locations (id, warehouse_id, code, name, is_default, is_active) VALUES (?, ?, 'MAIN', 'Main Location', 1, 1)", (loc_id, wh_id))
            loc_by_wh[wh_id] = loc_id

    # Categories
    categories_needed = {p['category'] for p in products if p['category']}
    c.execute("SELECT name, id FROM nbin_categories")
    existing_cats = dict(c.fetchall())
    for cat_name in categories_needed:
        if cat_name not in existing_cats:
            cat_id = str(uuid.uuid5(uuid.NAMESPACE_DNS, f"bishnoi.category.{cat_name}"))
            c.execute("INSERT INTO nbin_categories (id, name, is_active, created_at, updated_at) VALUES (?, ?, 1, ?, ?)", (cat_id, cat_name, now, now))
            existing_cats[cat_name] = cat_id

    # Brands
    c.execute("SELECT name, id FROM nbin_brands")
    existing_brands = dict(c.fetchall())
    for b_name in ['2MG', 'Getmeds']:
        if b_name not in existing_brands:
            b_id = str(uuid.uuid5(uuid.NAMESPACE_DNS, f"bishnoi.brand.{b_name}"))
            c.execute("INSERT INTO nbin_brands (id, name, is_active, created_at, updated_at) VALUES (?, ?, 1, ?, ?)", (b_id, b_name, now, now))
            existing_brands[b_name] = b_id

    # Units
    c.execute("SELECT code FROM nbin_units")
    existing_units = {r[0] for r in c.fetchall()}
    for p in products:
        uc = p['unit_code']
        if uc not in existing_units:
            c.execute("INSERT INTO nbin_units (code, name, is_active) VALUES (?, ?, 1)", (uc, p['unit_raw'].title()))
            existing_units.add(uc)

    # Products
    c.execute("SELECT sku, id FROM nbin_products")
    existing_prods = dict(c.fetchall())

    for p in products:
        sku = p['sku']
        name = p['brand']
        desc = f"{p['generic']} {p['strength']} {p['form']}".strip()
        cat_id = existing_cats.get(p['category'])
        brand_id = existing_brands.get(p['account'], existing_brands.get('Getmeds'))
        p_type = 'stock'
        u_code = p['unit_code']
        cost_p = p['dist_u_cents'] or p['dist_p_cents']
        sell_p = p['srp_u_cents'] or p['pat_u_cents'] or p['srp_p_cents']
        tax_bp = 1200 if p['is_vat'] else 0
        attrs = json.dumps({'account': p['account'], 'category': p['category'], 'generic_name': p['generic'], 'dosage_strength': p['strength']})

        if sku in existing_prods:
            db_id = existing_prods[sku]
            c.execute("""
                UPDATE nbin_products
                SET name = ?, description = ?, category_id = ?, brand_id = ?, product_type = ?,
                    unit_code = ?, cost_price = ?, selling_price = ?, tax_rate_bp = ?,
                    attributes = ?, status = 'active', updated_at = ?
                WHERE id = ?
            """, (name, desc, cat_id, brand_id, p_type, u_code, cost_p, sell_p, tax_bp, attrs, now, db_id))
        else:
            p_id = p['id']
            c.execute("""
                INSERT INTO nbin_products (
                    id, sku, name, description, category_id, brand_id, product_type,
                    unit_code, cost_price, selling_price, tax_rate_bp, tax_inclusive,
                    min_stock, reorder_point, reorder_qty, status, track_batches,
                    track_expiry, track_serials, attributes, created_at, updated_at
                ) VALUES (
                    ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 1,
                    0, 0, 0, 'active', 1, 1, 0, ?, ?, ?
                )
            """, (p_id, sku, name, desc, cat_id, brand_id, p_type, u_code, cost_p, sell_p, tax_bp, attrs, now, now))
            existing_prods[sku] = p_id

    # Stock 0
    c.execute("SELECT product_id, warehouse_id, location_id FROM nbin_stock")
    existing_stocks = {(r[0], r[1], r[2]) for r in c.fetchall()}

    for p in products:
        prod_id = existing_prods.get(p['sku'], p['id'])
        for wh_id, _ in warehouses:
            loc_id = loc_by_wh[wh_id]
            key = (prod_id, wh_id, loc_id)
            if key not in existing_stocks:
                stock_id = str(uuid.uuid4())
                c.execute("""
                    INSERT INTO nbin_stock (id, product_id, warehouse_id, location_id, on_hand, damaged, updated_at)
                    VALUES (?, ?, ?, ?, 0, 0, ?)
                """, (stock_id, prod_id, wh_id, loc_id, now))
                existing_stocks.add(key)

    conn.commit()
    conn.close()
    print(f"Inventory SQLite {db_name}: Synced {len(products)} products and zero stock.")

def sync_crm_sqlite(db_path, db_name, products, bundles, promos):
    print(f"\n--- Synchronizing CRM SQLite ({db_name}: {db_path}) ---")
    if not os.path.exists(db_path):
        print(f"Skipping {db_path} (does not exist)")
        return
    make_backup(db_path)
    now = datetime.now(UTC).strftime("%Y-%m-%d %H:%M:%S")

    conn = sqlite3.connect(db_path)
    c = conn.cursor()

    # Price Types
    c.execute("SELECT code FROM nbcm_price_types")
    existing_pts = {r[0] for r in c.fetchall()}
    needed_pts = [
        ('DOCTOR', "Doctor's Price", 1),
        ('PATIENT', "Patient Price", 2),
        ('SRP', "Suggested Retail Price (SRP)", 3),
        ('TRADE', "Drugstore / Hospital Trade Price", 4),
    ]
    for pt_code, pt_name, pt_sort in needed_pts:
        if pt_code not in existing_pts:
            c.execute("""
                INSERT INTO nbcm_price_types (code, name, entry, rule, requires_notes, needs_approval, is_active, sort_order)
                VALUES (?, ?, 'manual', 'custom', 0, 0, 1, ?)
            """, (pt_code, pt_name, pt_sort))
            existing_pts.add(pt_code)

    # Products
    c.execute("SELECT code, id FROM nbcm_products")
    existing_crm_prods = dict(c.fetchall())
    prod_code_to_id = {}

    for p in products:
        sku = p['sku']
        name = p['brand']
        gen = p['generic']
        st = p['strength']
        form = p['form']
        unit = p['unit_raw']
        vat_cat = 'vat' if p['is_vat'] else 'exempt'
        inv_id = p['id']
        extra_json = json.dumps({'account': p['account'], 'category': p['category'], 'pack_size': p['pack_size'], 'count': p['count']})

        if sku in existing_crm_prods:
            u_id = existing_crm_prods[sku]
            c.execute("""
                UPDATE nbcm_products
                SET name = ?, generic_name = ?, strength = ?, form = ?, unit = ?,
                    is_active = 1, vat_category = ?, inventory_id = ?, extra = ?, updated_at = ?
                WHERE id = ?
            """, (name, gen, st, form, unit, vat_cat, inv_id, extra_json, now, u_id))
            prod_code_to_id[sku] = u_id
        else:
            u_id = str(uuid.uuid4())
            c.execute("""
                INSERT INTO nbcm_products (id, code, name, generic_name, strength, form, unit, is_active, vat_category, inventory_id, extra, created_at, updated_at)
                VALUES (?, ?, ?, ?, ?, ?, ?, 1, ?, ?, ?, ?, ?)
            """, (u_id, sku, name, gen, st, form, unit, vat_cat, inv_id, extra_json, now, now))
            existing_crm_prods[sku] = u_id
            prod_code_to_id[sku] = u_id

        # Prices
        prices_to_set = [
            ('DOCTOR', p['doc_u_cents'] or p['doc_p_cents']),
            ('PATIENT', p['pat_u_cents'] or p['pat_p_cents']),
            ('SRP', p['srp_u_cents'] or p['srp_p_cents']),
            ('TRADE', p['trade_u_cents'] or p['trade_p_cents']),
        ]
        for pt_code, amt in prices_to_set:
            if amt > 0:
                c.execute("INSERT OR REPLACE INTO nbcm_product_prices (product_id, price_type_code, amount_centavos) VALUES (?, ?, ?)", (u_id, pt_code, amt))

    # Bundles
    c.execute("SELECT code, id FROM nbcm_bundles")
    existing_bundles = dict(c.fetchall())
    for b in bundles:
        b_code = b['code']
        b_name = b['name']
        b_price = b['price_cents']
        mapped_items = [{'productId': prod_code_to_id.get(it.get('sku')), 'name': it['name'], 'qty': it['qty']} for it in b['items']]
        items_json = json.dumps(mapped_items)

        if b_code in existing_bundles:
            b_id = existing_bundles[b_code]
            c.execute("UPDATE nbcm_bundles SET name = ?, price_centavos = ?, items = ?, is_active = 1 WHERE id = ?", (b_name, b_price, items_json, b_id))
        else:
            b_id = str(uuid.uuid4())
            c.execute("INSERT INTO nbcm_bundles (id, code, name, price_centavos, as_one_line, items, is_active, created_at, team_ids) VALUES (?, ?, ?, ?, 0, ?, 1, ?, '[]')", (b_id, b_code, b_name, b_price, items_json, now))

    # Promos
    c.execute("SELECT name, id FROM nbcm_promos")
    existing_promos = dict(c.fetchall())
    for pr in promos:
        pr_name = pr['name']
        pr_kind = pr['kind']
        pr_pct = pr['percent_bp']
        p_ids = [pr.get('product_id')] if pr.get('product_id') else []
        p_ids_json = json.dumps(p_ids)

        if pr_name in existing_promos:
            pr_id = existing_promos[pr_name]
            c.execute("UPDATE nbcm_promos SET kind = ?, percent_bp = ?, product_ids = ?, is_active = 1 WHERE id = ?", (pr_kind, pr_pct, p_ids_json, pr_id))
        else:
            pr_id = str(uuid.uuid4())
            c.execute("INSERT INTO nbcm_promos (id, name, kind, percent_bp, tiers, product_ids, is_active, created_at, team_ids) VALUES (?, ?, ?, ?, '[]', ?, 1, ?, '[]')", (pr_id, pr_name, pr_kind, pr_pct, p_ids_json, now))

    conn.commit()
    conn.close()
    print(f"CRM SQLite {db_name}: Synced {len(products)} products, {len(bundles)} bundles, and {len(promos)} promos.")

def main():
    print("=================================================================")
    print("   BISHNOI ONE INVENTORY & CRM PRICELIST POPULATION PIPELINE     ")
    print("=================================================================")

    products, bundles, promos = load_excel_data()
    print(f"Loaded {len(products)} products from sheet 'Pricelist02.2026'")
    print(f"Loaded {len(bundles)} bundles from sheet 'SEPT PROMO BUNDLES'")
    print(f"Loaded {len(promos)} promos from sheet 'B2C Promo'")

    # 1. Sync Inventory PostgreSQL (Production)
    sync_inventory_postgres(INV_GETMEDS_URL, "neondb_getmeds", products)
    sync_inventory_postgres(INV_BASE_URL, "neondb", products)

    # 2. Sync CRM PostgreSQL (Production)
    sync_crm_postgres(CRM_GETMEDS_URL, "neondb_getmeds", products, bundles, promos)
    sync_crm_postgres(CRM_BASE_URL, "neondb", products, bundles, promos)

    # 3. Sync Inventory SQLite
    sync_inventory_sqlite(INV_GETMEDS_SQLITE, "inventory-getmeds.db", products)
    sync_inventory_sqlite(INV_BASE_SQLITE, "inventory.db", products)

    # 4. Sync CRM SQLite
    sync_crm_sqlite(CRM_GETMEDS_SQLITE, "crm-getmeds.db", products, bundles, promos)
    sync_crm_sqlite(CRM_BASE_SQLITE, "crm.db", products, bundles, promos)

    print("\n=================================================================")
    print("             POPULATION PIPELINE COMPLETED SUCCESSFULLY          ")
    print("=================================================================")

if __name__ == '__main__':
    main()
