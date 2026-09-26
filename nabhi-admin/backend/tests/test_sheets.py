"""Excel reports and imports: permits, expatriates, documents, properties, office requests."""

import io
from datetime import UTC, datetime, timedelta

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app.config import Settings
from app.main import create_app
from app.models import User
from app.security import hash_password

PASSWORD = "testpassword123"


@pytest.fixture
def client(tmp_path):
    settings = Settings(
        _env_file=None,
        allowed_hosts=["testserver"],
        database_url=f"sqlite:///{(tmp_path / 'sheets.db').as_posix()}",
        org_timezone="Asia/Manila",
    )
    app = create_app(settings)
    with app.state.sessions() as db:
        db.add(User(email="admin@example.com", full_name="Admin Exec", password_hash=hash_password(PASSWORD), role="admin"))
        db.commit()
    with TestClient(app) as c:
        yield c


def sign_in(client, email="admin@example.com"):
    r = client.post("/auth/login", json={"email": email, "password": PASSWORD})
    assert r.status_code == 200
    return {"Authorization": f"Bearer {r.json()['accessToken']}"}


def day(n: int) -> str:
    return (datetime.now(UTC) + timedelta(days=n)).date().isoformat()


def xlsx(rows: list[list]) -> bytes:
    book = Workbook()
    for row in rows:
        book.active.append(row)
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def add_permit(client, h, ref="LGU-2026-001", days=10):
    r = client.post(
        "/admin/permits",
        headers=h,
        json={"title": "Mayor's permit", "category": "lgu", "agency": "LGU Makati", "referenceNumber": ref, "expirationDate": day(days)},
    )
    assert r.status_code == 201


def test_permit_report_has_days_left_and_standing(client):
    h = sign_in(client)
    add_permit(client, h, "LGU-1", days=10)
    add_permit(client, h, "LGU-2", days=-3)
    add_permit(client, h, "LGU-3", days=200)
    r = client.get("/admin/reports/permits", headers=h)
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert "Permits" in r.headers["content-disposition"]
    sheet = load_workbook(io.BytesIO(r.content)).active
    rows = list(sheet.iter_rows(values_only=True))
    head = rows[0]
    assert head[:3] == ("Title", "Category", "Agency")
    assert head[-2:] == ("Days left", "Standing")
    standing = {row[3]: row[-1] for row in rows[1:]}
    assert standing == {"LGU-2": "Overdue", "LGU-1": "Due soon", "LGU-3": "Current"}
    assert rows[1][1] == "Mayor's business permit"  # the label, not the stored code


def test_overview_report_has_summary_and_a_sheet_per_area(client):
    h = sign_in(client)
    add_permit(client, h, "LGU-1", days=-1)
    r = client.get("/admin/reports/overview", headers=h)
    assert r.status_code == 200
    book = load_workbook(io.BytesIO(r.content))
    assert book.sheetnames == ["Summary", "Permits", "Expatriates", "Documents", "Properties", "Office requests"]
    summary = list(book["Summary"].iter_rows(values_only=True))
    permits = next(row for row in summary if row and row[0] == "Permits")
    assert permits[1:5] == (1, 0, 0, 1)
    assert permits[-1] == "At risk"


def test_template_lists_what_each_column_takes(client):
    h = sign_in(client)
    r = client.get("/admin/imports/properties/template", headers=h)
    assert r.status_code == 200
    book = load_workbook(io.BytesIO(r.content))
    assert book.sheetnames == ["Properties", "How to fill in"]
    guide = {row[0]: row for row in book["How to fill in"].iter_rows(min_row=2, values_only=True)}
    assert guide["Lease end"][1] == "Yes"
    assert "Staff residence" in guide["Type"][2]


def test_import_checks_then_adds_and_updates(client):
    h = sign_in(client)
    add_permit(client, h, "BOC-7", days=100)
    file = xlsx([
        ["Permit name", "Type", "Issued by", "Ref no", "Expiry", "Cost"],
        ["Importer accreditation", "BOC", "Bureau of Customs", "BOC-7", datetime(2027, 3, 31), "12,500.00"],
        ["FDA license", "FDA license to operate", "FDA", "FDA-1", "31 Mar 2027", ""],
        ["No reference", "bir", "BIR", "", "2027-01-01", ""],
        ["Bad date", "bir", "BIR", "BIR-9", "soon", ""],
        ["Unknown kind", "parking", "BIR", "BIR-10", "2027-01-01", ""],
        ["FDA license again", "fda", "FDA", "FDA-1", "2027-04-01", ""],
    ])
    r = client.post("/admin/imports/permits/read", headers=h, files={"file": ("permits.xlsx", file, "application/octet-stream")})
    assert r.status_code == 200, r.text
    read = r.json()
    m = read["mapping"]
    assert (m["title"], m["category"], m["agency"], m["referenceNumber"], m["expirationDate"], m["estimatedCost"]) == (0, 1, 2, 3, 4, 5)
    assert read["rows"][0][4] == "2027-03-31"

    body = {"mapping": m, "rows": read["rows"], "dryRun": True}
    r = client.post("/admin/imports/permits", headers=h, json=body)
    assert r.status_code == 200, r.text
    check = r.json()
    assert (check["added"], check["updated"]) == (1, 1)
    problems = {p["row"]: p["message"] for p in check["problems"]}
    assert set(problems) == {4, 5, 6, 7}
    assert "Reference number is empty" in problems[4]
    assert "isn't a date" in problems[5]
    assert "isn't one of" in problems[6]
    assert "Same permit as row 3" in problems[7]
    assert "Status" in check["unmapped"]
    # nothing saved on a check
    assert len(client.get("/admin/permits", headers=h).json()) == 1

    r = client.post("/admin/imports/permits", headers=h, json={**body, "dryRun": False})
    assert (r.json()["added"], r.json()["updated"]) == (1, 1)
    permits = {p["referenceNumber"]: p for p in client.get("/admin/permits", headers=h).json()}
    assert set(permits) == {"BOC-7", "FDA-1"}
    assert permits["BOC-7"]["title"] == "Importer accreditation"
    assert permits["BOC-7"]["category"] == "boc"
    assert permits["BOC-7"]["estimatedCost"] == 12500.0
    assert permits["BOC-7"]["expirationDate"].startswith("2027-03-31")
    assert permits["FDA-1"]["status"] == "active"


def test_report_imports_back_unchanged(client):
    h = sign_in(client)
    r = client.post(
        "/admin/properties",
        headers=h,
        json={
            "name": "Makati head office", "propertyType": "expat_housing", "address": "Ayala Ave", "lessorName": "Ayala Land",
            "leaseEnd": day(300), "monthlyRent": 250000, "status": "renewal_negotiation",
        },
    )
    assert r.status_code == 201
    report = client.get("/admin/reports/properties", headers=h).content
    read = client.post("/admin/imports/properties/read", headers=h, files={"file": ("p.xlsx", report, "application/octet-stream")}).json()
    assert None not in read["mapping"].values()
    r = client.post("/admin/imports/properties", headers=h, json={"mapping": read["mapping"], "rows": read["rows"], "dryRun": False})
    assert r.json() == {"added": 0, "updated": 1, "problems": [], "unmapped": []}
    (prop,) = client.get("/admin/properties", headers=h).json()
    assert (prop["propertyType"], prop["status"], prop["monthlyRent"]) == ("expat_housing", "renewal_negotiation", 250000.0)


def test_csv_import_of_expatriates(client):
    h = sign_in(client)
    csv = "Name,Position,Nationality,Document,Document number,Expires\nKenji Sato,Country manager,Japanese,Visa,9G-100,2027-06-30\n"
    read = client.post("/admin/imports/expatriates/read", headers=h, files={"file": ("x.csv", csv.encode(), "text/csv")}).json()
    r = client.post("/admin/imports/expatriates", headers=h, json={"mapping": read["mapping"], "rows": read["rows"], "dryRun": False})
    assert r.json()["added"] == 1
    (rec,) = client.get("/admin/immigration", headers=h).json()
    assert (rec["individualName"], rec["documentType"]) == ("Kenji Sato", "visa")


def test_import_needs_the_required_columns(client):
    h = sign_in(client)
    r = client.post("/admin/imports/documents", headers=h, json={"mapping": {"title": 0}, "rows": [["x"]], "dryRun": True})
    assert r.status_code == 422
    assert "Category" in r.json()["detail"]


def test_requests_are_export_only(client):
    h = sign_in(client)
    assert client.get("/admin/reports/requests", headers=h).status_code == 200
    assert client.post("/admin/imports/requests", headers=h, json={"mapping": {}, "rows": []}).status_code == 422


def test_view_only_role_can_export_but_not_import(client):
    from app.models import User as U
    from app.permissions import ROLE_AUDITOR_ID

    with client.app.state.sessions() as db:
        db.add(U(email="audit@example.com", full_name="Auditor", password_hash=hash_password(PASSWORD), role_id=ROLE_AUDITOR_ID))
        db.commit()
    h = sign_in(client, "audit@example.com")
    assert client.get("/admin/reports/permits", headers=h).status_code == 200
    assert client.get("/admin/reports/overview", headers=h).status_code == 200
    assert client.get("/admin/imports/permits/template", headers=h).status_code == 403
    r = client.post("/admin/imports/permits", headers=h, json={"mapping": {}, "rows": [], "dryRun": True})
    assert r.status_code == 403
