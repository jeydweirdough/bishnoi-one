"""Fleet, travel and devices: the executive assistant's trackers (app/routers/operations.py),
their countdowns on the overview and alerts, and the import from the JASMINE TASKS workbook."""

from datetime import date, datetime, timedelta

from openpyxl import Workbook

from app.models import Notification, User
from app.permissions import ROLE_AUDITOR_ID
from app.security import hash_password
from app.tools import import_jasmine
from tests.test_rbac import PASSWORD, app, client, sign_in  # noqa: F401 (fixtures)


def soon(days: int) -> str:
    return (date.today() + timedelta(days=days)).isoformat()


def test_vehicles_count_down_and_raise_alerts_that_clear_themselves(app, client):
    asst = sign_in(client, "asst@example.com")  # the Executive Assistant adds and changes them
    r = client.post("/admin/vehicles", json={"name": "Ford Ranger AT", "plate": "NFY 4260", "registrationDue": soon(10), "insuranceDue": soon(200), "pmsDue": soon(-3)}, headers=asst)
    assert r.status_code == 201, r.text
    car = r.json()
    assert [d["what"] for d in car["dues"]] == ["pms", "registration", "insurance"]  # soonest first
    assert (car["nextDue"], car["daysRemaining"], car["urgencyTier"]) == ("Maintenance (PMS)", -3, "overdue")

    with app.state.sessions() as db:
        open_alerts = {n.entity_type for n in db.query(Notification).filter_by(entity_id=car["id"], is_resolved=False)}
    assert open_alerts == {"vehicle_pms", "vehicle_registration"}  # insurance is too far off

    dash = client.get("/admin/dashboard", headers=sign_in(client)).json()
    assert {h["entityType"] for h in dash["countdownHighlights"]} >= {"vehicle_pms", "vehicle_registration"}
    assert next(a for a in dash["health"] if a["key"] == "vehicles")["overdue"] == 1
    assert dash["stats"]["vehicles"] == {"total": 1, "dueSoon": 1, "overdue": 1}

    # Maintenance done and the next one set far ahead: its alert closes by itself.
    r = client.patch(f"/admin/vehicles/{car['id']}", json={**{k: car[k] for k in ("name", "plate", "registrationDue", "insuranceDue")}, "pmsDue": soon(180)}, headers=asst)
    assert r.status_code == 200, r.text
    with app.state.sessions() as db:
        assert {n.entity_type for n in db.query(Notification).filter_by(entity_id=car["id"], is_resolved=False)} == {"vehicle_registration"}

    assert client.delete(f"/admin/vehicles/{car['id']}", headers=asst).status_code == 403  # the assistant doesn't delete
    assert client.delete(f"/admin/vehicles/{car['id']}", headers=sign_in(client)).status_code == 204
    with app.state.sessions() as db:
        assert db.query(Notification).filter_by(entity_id=car["id"]).count() == 0


def test_cards_trips_and_devices_are_checked_and_logged(app, client):
    asst = sign_in(client, "asst@example.com")
    card = client.post("/admin/fleet-cards", json={"holderName": "Annie Maglatang", "monthlyAllowance": 15000, "plate": "NJG-6869", "status": "handed_over"}, headers=asst)
    assert card.status_code == 201 and card.json()["cardState"] == "active"
    assert client.post("/admin/fleet-cards", json={"holderName": "X", "status": "lost"}, headers=asst).status_code == 422

    trip = client.post("/admin/travel", json={"repName": "Paolo Paradero", "doctorName": "Dr. Justine Natividad", "kind": "hotel", "startsOn": soon(5), "travelDates": "in 5 days", "status": "for_approval"}, headers=asst).json()
    assert trip["daysRemaining"] == 5
    dash = client.get("/admin/dashboard", headers=sign_in(client)).json()
    assert any(h["entityType"] == "travel" and h["title"] == "Dr. Justine Natividad" for h in dash["countdownHighlights"])  # close and not approved
    assert dash["stats"]["travel"]["waiting"] == 1

    assert client.post("/admin/devices", json={"issuedTo": "Jade Madigal"}, headers=asst).status_code == 422  # nothing handed over
    device = client.post("/admin/devices", json={"issuedTo": "Jade Madigal", "phone": "Galaxy A06 5G", "condition": "brand_new"}, headers=asst)
    assert device.status_code == 201

    log = client.get("/admin/audit", headers=sign_in(client)).json()
    actions = {e["action"] for e in (log["items"] if isinstance(log, dict) else log)}
    assert {"fleet-cards.added", "travel.added", "devices.added"} <= actions


def test_the_auditor_sees_but_cannot_change(app, client):
    with app.state.sessions() as db:
        db.add(User(email="audit@example.com", full_name="Auditor", password_hash=hash_password(PASSWORD), role_id=ROLE_AUDITOR_ID))
        db.commit()
    auditor = sign_in(client, "audit@example.com")
    assert client.get("/admin/vehicles", headers=auditor).status_code == 200
    assert client.post("/admin/vehicles", json={"name": "Kia Soluto"}, headers=auditor).status_code == 403


def test_the_workbook_comes_in_cleaned_with_nothing_lost(tmp_path):
    book = Workbook()
    book.active.title = "VEHICLES MONITORING"
    book.active.append(["CARS NAME", "PLATE NUMBER", "Year Model", "Color", "OWNER'S NAME", "RENEWAL DATE", None, "INSURANCE", None, "PMS Date", "Fleet card", "Parts"])
    book.active.append(["FORD RANGER MT", "NDA 8790", 2019, "BLACK", "2MG", "October 31, 2025", "October 31, 2026", datetime(2025, 7, 28), "with BF sticker", "PMS", None, None])
    book.active.append(["LIST OF INSURANCE RENEWAL", None, None, None, None, None, None, None, None, None, None, None])
    s = book.create_sheet("SALES PERSON_FLEETCARD ")
    s.append(["DATE", "NAME", "COMPANY", "ALLOCATION", "FUEL", "PLATE", "STICKER", "UNIT", "STATUS", "REMARKS"])
    s.append([datetime(2026, 1, 5), "ANNIE MAGLATANG", "2MG", "200 LITERS MONTHLY", "GASOLINE", None, "ST 107B", "Veloz", "HANDED OVER TO HR", "BLOCKED"])
    s = book.create_sheet("FLIGHT AND HOTEL SPONSORSHIP")
    s.append(["REQUESTED", "MSA", "DOCTOR", "BIRTHDATE", "CHECK IN-OUT", "HOTEL", "FLIGHT", "TYPE", "STATUS", "REMARKS", "APPROVED"])
    s.append([datetime(2026, 9, 12), "ANDI BERMIDO", "DR. LEO BLANCO", datetime(1980, 1, 2), None, None, "November 24-27 (CEBPAC)", "FLIGHT BOOKING", "BOOKED", "DONE", None])
    s = book.create_sheet("LAPTOPMOBILE PHONE MONITORING ")
    s.append(["", "LAPTOP", "QTY", "PHONE", "QTY", "SIM", "QTY", "PERSON", "REMARKS", "Signed by"])
    s.append([datetime(2026, 3, 25), "HP LAPTOP", 1, None, None, "DITO", 2, "C/O HR", "Old", "MICHAEL UY"])
    s.append(["June 8, 2026 - Received 7 pcs phone", None, None, None, None, None, None, None, None, None])
    path = tmp_path / "jasmine.xlsx"
    book.save(path)

    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    [car] = import_jasmine.vehicles(wb)  # the note under the list isn't a car
    assert (car.registration_due, car.insurance_due, car.pms_due, car.year_model) == (date(2026, 10, 31), date(2025, 7, 28), None, 2019)
    assert "with BF sticker" in car.notes and "PMS 'PMS'" in car.notes  # kept, word for word
    [card] = import_jasmine.fleet_cards(wb)
    assert (card.holder_name, card.monthly_allowance, card.allowance_note, card.conduction_sticker, card.status, card.card_state) == (
        "Annie Maglatang", None, "200 liters monthly", "ST 107B", "with_hr", "blocked")
    [trip] = import_jasmine.sponsorships(wb)
    assert (trip.kind, trip.doctor_name, trip.starts_on, trip.hotel_or_flight, trip.doctor_birthdate) == ("flight", "Dr. Leo Blanco", date(2026, 11, 24), "CEBPAC", date(1980, 1, 2))
    assert "taken as 2026" in trip.notes
    devices, stock = import_jasmine.devices(wb)
    assert [(d.laptop, d.sim, d.condition, d.signed_by) for d in devices] == [("HP LAPTOP", "DITO ×2", "used", "Michael Uy")]
    assert stock == ["June 8, 2026 - Received 7 pcs phone"]


def test_quick_steps_roll_dates_forward_and_build_the_history(app, client):
    asst = sign_in(client, "asst@example.com")
    car = client.post("/admin/vehicles", json={"name": "Kia Soluto", "plate": "NIQ 2591", "registrationDue": soon(-20), "pmsDue": soon(-300)}, headers=asst).json()

    r = client.post(f"/admin/vehicles/{car['id']}/log", json={"action": "registration_renewed"}, headers=asst)
    assert r.status_code == 200, r.text
    got = r.json()
    assert got["record"]["registrationDue"] == (date.fromisoformat(soon(-20)).replace(year=date.today().year + 1)).isoformat() or got["record"]["registrationDue"] > soon(0)
    r = client.post(f"/admin/vehicles/{car['id']}/log", json={"action": "pms_done"}, headers=asst).json()
    assert soon(170) <= r["record"]["pmsDue"] <= soon(190)  # six months from today
    assert r["record"]["urgencyTier"] != "overdue"
    client.post(f"/admin/vehicles/{car['id']}/log", json={"action": "note", "note": "Brake pads replaced"}, headers=asst)
    one = client.get(f"/admin/vehicles/{car['id']}", headers=asst).json()
    assert [e["kind"] for e in one["history"]] == ["note", "pms_done", "registration_renewed", "added"]
    assert one["history"][0]["by"] == "Assistant Exec" and "pms_done" in one["steps"]

    assert client.post(f"/admin/vehicles/{car['id']}/log", json={"action": "booked"}, headers=asst).status_code == 422  # not a car's step
    assert client.post(f"/admin/vehicles/{car['id']}/log", json={"action": "note"}, headers=asst).status_code == 422  # a note needs words

    trip = client.post("/admin/travel", json={"repName": "Paolo", "doctorName": "Dr. Lao", "kind": "flight", "status": "for_processing"}, headers=asst).json()
    booked = client.post(f"/admin/travel/{trip['id']}/log", json={"action": "booked"}, headers=asst).json()["record"]
    assert booked["status"] == "booked" and booked["approvedOn"] == date.today().isoformat()
