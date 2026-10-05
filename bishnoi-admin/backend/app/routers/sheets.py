"""Excel reports and imports for permits, expatriates, documents, properties
and office requests.

Reports: /admin/reports/{kind} is one list as an Excel workbook, with the days
left and whether each item is current, due soon or overdue. /admin/reports/overview
is one workbook with a summary sheet and a sheet per area the person may see.

Imports, in the CRM's three steps:
1. /admin/imports/{kind}/read: the file is read, never kept. What comes back is
   its headers, its rows, the fields they can go into, and a first guess at
   which column goes where (from the headers' names).
2. /admin/imports/{kind} with dryRun: every row is checked, nothing is saved.
3. /admin/imports/{kind}: the rows that pass are saved; the others are listed
   with the reason.

A report's columns are the import's fields, so a report can be edited in Excel
and imported back. An existing record is matched, not added twice: a permit by
its reference number, an expatriate document by its number, a document by its
reference number (or its title and counterparty when it has none), a property
by its name. /admin/imports/{kind}/template is an empty sheet with the right
headers and a second sheet saying what each column takes.

Reports need <area>.view; imports need <area>.create and <area>.edit.
Office requests are exported only.
"""

import csv
import io
import re
from collections.abc import Callable
from dataclasses import dataclass
from datetime import date, datetime
from typing import Annotated, Any, Literal

from fastapi import APIRouter, File, HTTPException, Response, UploadFile, status
from pydantic import BaseModel, Field, ValidationError
from sqlalchemy import select

from app.access import CurrentAccess
from app.api import Api
from app.audit import audit
from app.config import Settings
from app.deps import DB, SettingsDep
from app.executive import sync_notifications, today_in_org
from app.models import ComplianceRecord, CorporateDocument, ImmigrationRecord, PropertyRecord, Ticket, new_id, utcnow
from app.routers.admin import DocumentInput, ImmigrationInput, PermitInput, PropertyInput, parse_iso_datetime

router = APIRouter(prefix="/admin", tags=["reports"])

MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 5000
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


# --- what each list holds --------------------------------------------------------------------


@dataclass(frozen=True)
class Col:
    key: str  # the import field (the input's name); "" for report-only columns
    label: str  # the header, in reports and templates
    attr: str = ""  # the record's attribute
    type: Literal["text", "date", "int", "money", "percent", "bool", "choice"] = "text"
    required: bool = False
    choices: dict[str, str] | None = None  # stored value -> label
    aliases: tuple[str, ...] = ()
    get: Callable[[Any], Any] | None = None  # report-only value
    width: int = 18


@dataclass(frozen=True)
class Area:
    kind: str
    title: str  # sheet name and file name
    area: str  # permission area
    model: type
    columns: tuple[Col, ...]
    due: Callable[[Any], datetime | None]  # the countdown date
    window: int  # days before the due date that count as "due soon" (as on the overview)
    input: type[BaseModel] | None = None  # None: report only
    order: Any = None
    noun: str = "record"
    match: Callable[[Any], Any] | None = None  # the input -> a where clause for the same record
    ident: Callable[[Any], str] | None = None  # the input -> what makes it the same record, to catch repeats in a file


PERMIT_CATEGORIES = {
    "boc": "Bureau of Customs (BOC)",
    "lgu": "Mayor's business permit",
    "sec": "SEC GIS and filings",
    "bir": "BIR registration",
    "fda": "FDA license to operate",
    "sanitary_fire": "Fire safety and sanitary",
    "other": "Other",
}
PERMIT_STATUSES = {"active": "Active", "pending_renewal": "For renewal", "under_review": "Under agency review", "expired": "Expired"}
EXPAT_TYPES = {
    "acr_card": "ACR I-Card",
    "aep_permit": "DOLE work permit (AEP)",
    "visa": "Visa",
    "passport": "Passport",
    "exec_id": "Company ID",
}
EXPAT_STATUSES = {"valid": "Valid", "expiring_soon": "Expiring soon", "in_renewal": "Renewal filed", "expired": "Expired"}
DOCUMENT_CATEGORIES = {
    "insurance": "Insurance policy",
    "commercial_contract": "Contract",
    "contract": "Contract",
    "nda": "NDA",
    "board_resolution": "Board resolution",
    "legal_retainer": "Legal or consultant retainer",
    "sla": "Service level agreement",
}
DOCUMENT_STATUSES = {"active": "Active", "in_review": "In review", "renewed": "Renewed", "terminated": "Terminated"}
PROPERTY_TYPES = {"office": "Office", "warehouse": "Warehouse", "depot": "Depot", "parking": "Parking", "expat_housing": "Staff residence"}
PROPERTY_STATUSES = {
    "active": "Active",
    "renewal_negotiation": "Under negotiation",
    "expiring_soon": "Expiring soon",
    "terminated": "Terminated",
}
REQUEST_CATEGORIES = {
    "facilities": "Facilities and repairs",
    "vip_logistics": "Guest transport",
    "supplies": "Office supplies",
    "boardroom": "Meeting rooms",
    "courier_dispatch": "Courier and government liaison",
    "it_security": "IT equipment and access cards",
    "general": "Other",
}
REQUEST_STATUSES = {
    "open": "Open",
    "in_progress": "In progress",
    "pending_vendor": "Waiting on supplier",
    "resolved": "Resolved",
    "closed": "Closed",
}
PRIORITIES = {"urgent": "Urgent", "high": "High", "normal": "Normal", "low": "Low"}


AREAS: dict[str, Area] = {
    "permits": Area(
        kind="permits",
        title="Permits",
        area="permits",
        model=ComplianceRecord,
        input=PermitInput,
        noun="permit",
        due=lambda r: r.expiration_date,
        window=30,
        order=ComplianceRecord.expiration_date.asc(),
        match=lambda p: ComplianceRecord.reference_number == p.referenceNumber.strip(),
        ident=lambda p: normal(p.referenceNumber),
        columns=(
            Col("title", "Title", "title", required=True, aliases=("permit", "name", "permitname"), width=34),
            Col("category", "Category", "category", "choice", True, PERMIT_CATEGORIES, ("type",), width=26),
            Col("agency", "Agency", "agency", required=True, aliases=("issuedby", "issuer"), width=26),
            Col("referenceNumber", "Reference number", "reference_number", required=True, aliases=("reference", "refno", "permitno", "number")),
            Col("issuanceDate", "Issued", "issuance_date", "date", aliases=("issuancedate", "dateissued", "issuedon")),
            Col("expirationDate", "Expires", "expiration_date", "date", True, aliases=("expirationdate", "expiry", "expirydate", "validuntil", "duedate")),
            Col("renewalWindowDays", "Renewal window (days)", "renewal_window_days", "int", aliases=("renewalwindow", "window", "leaddays")),
            Col("status", "Status", "status", "choice", choices=PERMIT_STATUSES),
            Col("responsiblePerson", "Responsible person", "responsible_person", aliases=("responsible", "owner", "incharge")),
            Col("estimatedCost", "Estimated cost", "estimated_cost", "money", aliases=("cost", "renewalcost", "fee", "fees")),
            Col("notes", "Notes", "notes", aliases=("remarks", "note"), width=40),
        ),
    ),
    "expatriates": Area(
        kind="expatriates",
        title="Expatriates",
        area="immigration",
        model=ImmigrationRecord,
        input=ImmigrationInput,
        noun="expatriate document",
        due=lambda r: r.expiration_date,
        window=60,
        order=ImmigrationRecord.expiration_date.asc(),
        match=lambda p: ImmigrationRecord.document_number == p.documentNumber.strip(),
        ident=lambda p: normal(p.documentNumber),
        columns=(
            Col("individualName", "Name", "individual_name", required=True, aliases=("fullname", "employee", "expatriate", "individualname"), width=26),
            Col("designation", "Position", "designation", required=True, aliases=("designation", "title", "jobtitle")),
            Col("nationality", "Nationality", "nationality", required=True, aliases=("country", "citizenship")),
            Col("documentType", "Document", "document_type", "choice", True, EXPAT_TYPES, ("documenttype", "type"), width=24),
            Col("documentNumber", "Document number", "document_number", required=True, aliases=("number", "docno", "documentno")),
            Col("issuanceDate", "Issued", "issuance_date", "date", aliases=("issuancedate", "dateissued")),
            Col("expirationDate", "Expires", "expiration_date", "date", True, aliases=("expirationdate", "expiry", "expirydate", "validuntil")),
            Col("renewalLeadDays", "Renewal lead time (days)", "renewal_lead_days", "int", aliases=("renewallead", "leaddays", "leadtime")),
            Col("sponsoringEntity", "Sponsor", "sponsoring_entity", aliases=("sponsoringentity", "sponsoringcompany", "employer")),
            Col("status", "Status", "status", "choice", choices=EXPAT_STATUSES),
            Col("assignedAssistant", "Handled by", "assigned_assistant", aliases=("assignedassistant", "assistant", "assignedto")),
            Col("notes", "Notes", "notes", aliases=("remarks", "note"), width=40),
        ),
    ),
    "documents": Area(
        kind="documents",
        title="Documents",
        area="documents",
        model=CorporateDocument,
        input=DocumentInput,
        noun="document",
        due=lambda r: r.expiration_date,
        window=60,
        order=CorporateDocument.expiration_date.asc(),
        match=lambda p: (
            CorporateDocument.reference_number == p.referenceNumber.strip()
            if p.referenceNumber
            else (CorporateDocument.title == p.title.strip()) & (CorporateDocument.counterparty == p.counterparty.strip())
        ),
        ident=lambda p: normal(p.referenceNumber) if p.referenceNumber else f"{normal(p.title)}|{normal(p.counterparty)}",
        columns=(
            Col("title", "Title", "title", required=True, aliases=("document", "name", "documentname"), width=34),
            Col("category", "Category", "category", "choice", True, DOCUMENT_CATEGORIES, ("type", "documenttype"), width=24),
            Col("counterparty", "Counterparty", "counterparty", required=True, aliases=("party", "vendor", "supplier", "insurer", "company"), width=26),
            Col("referenceNumber", "Reference number", "reference_number", aliases=("reference", "refno", "policyno", "contractno", "number")),
            Col("effectiveDate", "Effective", "effective_date", "date", aliases=("effectivedate", "startdate", "start")),
            Col("expirationDate", "Expires", "expiration_date", "date", aliases=("expirationdate", "expiry", "enddate", "end")),
            Col("autoRenews", "Renews automatically", "auto_renews", "bool", aliases=("autorenews", "autorenew", "autorenewal")),
            Col("noticePeriodDays", "Notice period (days)", "notice_period_days", "int", aliases=("noticeperiod", "notice", "noticedays")),
            Col("custodian", "Custodian", "custodian", aliases=("keptby", "owner")),
            Col("status", "Status", "status", "choice", choices=DOCUMENT_STATUSES),
            Col("annualValue", "Annual value", "annual_value", "money", aliases=("value", "annualcost", "premium", "amount")),
            Col("notes", "Notes", "notes", aliases=("remarks", "note"), width=40),
        ),
    ),
    "properties": Area(
        kind="properties",
        title="Properties",
        area="properties",
        model=PropertyRecord,
        input=PropertyInput,
        noun="property",
        due=lambda r: r.lease_end,
        window=90,
        order=PropertyRecord.lease_end.asc(),
        match=lambda p: PropertyRecord.name == p.name.strip(),
        ident=lambda p: normal(p.name),
        columns=(
            Col("name", "Property", "name", required=True, aliases=("name", "propertyname", "site"), width=28),
            Col("propertyType", "Type", "property_type", "choice", True, PROPERTY_TYPES, ("propertytype",)),
            Col("address", "Address", "address", required=True, aliases=("location",), width=36),
            Col("lessorName", "Lessor", "lessor_name", required=True, aliases=("lessorname", "landlord", "owner"), width=24),
            Col("lessorContact", "Lessor contact", "lessor_contact", aliases=("lessorcontact", "landlordcontact", "contact")),
            Col("leaseStart", "Lease start", "lease_start", "date", aliases=("leasestart", "start", "startdate")),
            Col("leaseEnd", "Lease end", "lease_end", "date", True, aliases=("leaseend", "end", "enddate", "expiry")),
            Col("noticeDeadlineDays", "Notice period (days)", "notice_deadline_days", "int", aliases=("noticedeadline", "notice", "noticedays")),
            Col("monthlyRent", "Monthly rent", "monthly_rent", "money", aliases=("rent", "monthlyrent")),
            Col("escalationRatePct", "Yearly increase (%)", "escalation_rate_pct", "percent", aliases=("escalation", "escalationrate", "increase")),
            Col("securityDeposit", "Security deposit", "security_deposit", "money", aliases=("deposit", "securitydeposit")),
            Col("utilitiesInfo", "Utility accounts", "utilities_info", aliases=("utilities", "utilitiesinfo"), width=30),
            Col("status", "Status", "status", "choice", choices=PROPERTY_STATUSES),
            Col("assignedManager", "Managed by", "assigned_manager", aliases=("assignedmanager", "manager")),
            Col("notes", "Notes", "notes", aliases=("remarks", "note"), width=40),
        ),
    ),
    "requests": Area(
        kind="requests",
        title="Office requests",
        area="tickets",
        model=Ticket,
        due=lambda r: r.sla_due_date if r.status in ("open", "in_progress", "pending_vendor") else None,
        window=3,
        order=Ticket.created_at.desc(),
        columns=(
            Col("", "Request number", get=lambda t: t.ticket_number, width=16),
            Col("", "Title", get=lambda t: t.title, width=34),
            Col("", "Category", "category", "choice", choices=REQUEST_CATEGORIES, width=26),
            Col("", "Priority", "priority", "choice", choices=PRIORITIES, width=12),
            Col("", "Status", "status", "choice", choices=REQUEST_STATUSES),
            Col("", "Requested by", get=lambda t: t.requester_name),
            Col("", "Assigned to", get=lambda t: t.assigned_to),
            Col("", "Requested on", "created_at", "date"),
            Col("", "Target date", "sla_due_date", "date"),
            Col("", "Resolved on", "resolved_at", "date"),
            Col("", "Description", get=lambda t: t.description, width=40),
            Col("", "Resolution", get=lambda t: t.resolution_notes, width=40),
        ),
    ),
}

Kind = Literal["permits", "expatriates", "documents", "properties", "requests"]
ImportKind = Literal["permits", "expatriates", "documents", "properties"]


def normal(text: str) -> str:
    """A header or a value, as compared: lower case, letters and digits only."""
    return re.sub(r"[^a-z0-9]", "", str(text).lower())


# --- reports ---------------------------------------------------------------------------------


def standing(days: int | None, window: int, open_: bool = True) -> str:
    if days is None or not open_:
        return ""
    if days < 0:
        return "Overdue"
    if days <= window:
        return "Due soon"
    return "Current"


def as_date(value: datetime | None) -> date | None:
    return value.date() if isinstance(value, datetime) else value


def write_sheet(book, spec: Area, rows: list, today: date, first: bool = False) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    sheet = book.active if first else book.create_sheet()
    sheet.title = spec.title[:31]
    columns = [*spec.columns]
    heads = [c.label for c in columns] + ["Days left", "Standing"]
    sheet.append(heads)
    for rec in rows:
        values: list[Any] = []
        for c in columns:
            raw = c.get(rec) if c.get else getattr(rec, c.attr)
            if c.type == "date":
                raw = as_date(raw)
            elif c.type == "choice" and raw is not None:
                raw = (c.choices or {}).get(raw, str(raw).replace("_", " ").capitalize())
            elif c.type == "bool":
                raw = "Yes" if raw else "No"
            values.append(raw)
        due = as_date(spec.due(rec))
        days = (due - today).days if due else None
        values += [days, standing(days, spec.window)]
        sheet.append(values)

    bold = Font(bold=True)
    fill = PatternFill("solid", fgColor="EEF0F4")
    for i, head in enumerate(heads, start=1):
        top = sheet.cell(row=1, column=i)
        top.font, top.fill = bold, fill
        top.alignment = Alignment(vertical="center")
        col = columns[i - 1] if i <= len(columns) else None
        letter = get_column_letter(i)
        sheet.column_dimensions[letter].width = col.width if col else 12
        fmt = {"date": "yyyy-mm-dd", "money": "#,##0.00", "percent": "0.00"}.get(col.type if col else "", None)
        if fmt:
            for cell in sheet[letter][1:]:
                cell.number_format = fmt
    red = Font(color="B42318", bold=True)
    amber = Font(color="8A5A12", bold=True)
    for cell in sheet[get_column_letter(len(heads))][1:]:
        if cell.value == "Overdue":
            cell.font = red
        elif cell.value == "Due soon":
            cell.font = amber
    sheet.freeze_panes = "A2"
    if rows:
        sheet.auto_filter.ref = sheet.dimensions


def records(db, spec: Area) -> list:
    return list(db.scalars(select(spec.model).order_by(spec.order)).all())


def download(book, name: str) -> Response:
    out = io.BytesIO()
    book.save(out)
    return Response(
        out.getvalue(),
        media_type=XLSX,
        headers={"Content-Disposition": f'attachment; filename="{name}"', "Cache-Control": "no-store"},
    )


def stamp(settings: Settings) -> str:
    return today_in_org(settings).isoformat()


@router.get("/reports/overview")
def overview_report(access: CurrentAccess, db: DB, settings: SettingsDep) -> Response:
    """Summary sheet (each area against its target) and a sheet per area."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    shown = [s for s in AREAS.values() if access.can(f"{s.area}.view")]
    if not shown:
        raise HTTPException(status.HTTP_403_FORBIDDEN, "Your role doesn't allow that.")
    today = today_in_org(settings)
    book = Workbook()
    summary = book.active
    summary.title = "Summary"
    summary.append([f"Compliance report as of {today.isoformat()}"])
    summary["A1"].font = Font(bold=True, size=13)
    summary.append([])
    summary.append(["Area", "Total", "Current", "Due soon", "Overdue", "Due soon means", "Standing"])
    for cell in summary[3]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="EEF0F4")
    lists = {s.kind: records(db, s) for s in shown}
    for s in shown:
        dues = [as_date(s.due(r)) for r in lists[s.kind]]
        dues = [d for d in dues if d]
        overdue = sum(1 for d in dues if d < today)
        soon = sum(1 for d in dues if 0 <= (d - today).days <= s.window)
        if s.kind == "requests":
            opened = [r for r in lists[s.kind] if r.status in ("open", "in_progress", "pending_vendor")]
            total, label = len(opened), "Open office requests"
            note = f"Target date within {s.window} days"
        else:
            total, label = len(dues), s.title
            note = f"Expires within {s.window} days"
        word = "At risk" if overdue else "Needs attention" if total and soon / total > 0.25 else "On track"
        summary.append([label, total, total - overdue - soon, soon, overdue, note, word])
    for col, width in zip("ABCDEFG", (24, 10, 10, 10, 10, 30, 18), strict=True):
        summary.column_dimensions[col].width = width
    for s in shown:
        write_sheet(book, s, lists[s.kind], today)
    return download(book, f"Compliance report {stamp(settings)}.xlsx")


@router.get("/reports/{kind}")
def area_report(kind: Kind, access: CurrentAccess, db: DB, settings: SettingsDep) -> Response:
    from openpyxl import Workbook

    spec = AREAS[kind]
    access.require(f"{spec.area}.view")
    book = Workbook()
    write_sheet(book, spec, records(db, spec), today_in_org(settings), first=True)
    return download(book, f"{spec.title} {stamp(settings)}.xlsx")


@router.get("/imports/{kind}/template")
def import_template(kind: ImportKind, access: CurrentAccess) -> Response:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    spec = AREAS[kind]
    access.require(f"{spec.area}.create")
    book = Workbook()
    sheet = book.active
    sheet.title = spec.title
    sheet.append([c.label for c in spec.columns])
    for i, c in enumerate(spec.columns, start=1):
        top = sheet.cell(row=1, column=i)
        top.font = Font(bold=True)
        top.fill = PatternFill("solid", fgColor="EEF0F4")
        sheet.column_dimensions[top.column_letter].width = c.width
    sheet.freeze_panes = "A2"
    guide = book.create_sheet("How to fill in")
    guide.append(["Column", "Required", "What to write"])
    for cell in guide[1]:
        cell.font = Font(bold=True)
    for c in spec.columns:
        guide.append([c.label, "Yes" if c.required else "No", explain(c) or "Text"])
    guide.column_dimensions["A"].width = 26
    guide.column_dimensions["B"].width = 10
    guide.column_dimensions["C"].width = 80
    return download(book, f"{spec.title} import template.xlsx")


def explain(c: Col) -> str:
    if c.type == "choice":
        return "One of: " + ", ".join(dict.fromkeys((c.choices or {}).values()))
    return {
        "date": "A date, e.g. 2027-03-31 or 31 Mar 2027",
        "int": "A whole number of days",
        "money": "An amount in pesos, e.g. 12,500.00",
        "percent": "A number, e.g. 5 for 5%",
        "bool": "Yes or No",
    }.get(c.type, "")


# --- reading a file --------------------------------------------------------------------------


def cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def read_rows(name: str, data: bytes) -> list[list[str]]:
    if name.lower().endswith(".csv"):
        text = data.decode("utf-8-sig", errors="replace")
        return [[c.strip() for c in row] for row in csv.reader(io.StringIO(text))]
    if name.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        try:
            book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception:
            raise HTTPException(status.HTTP_422_UNPROCESSABLE_CONTENT, "That file couldn't be read as an Excel workbook.") from None
        sheet = book.worksheets[0]
        return [[cell(v) for v in row] for row in sheet.iter_rows(values_only=True)]
    raise HTTPException(status.HTTP_422_UNPROCESSABLE_CONTENT, "Use an Excel file (.xlsx) or a CSV file (.csv).")


def fields_of(spec: Area) -> list[Col]:
    return [c for c in spec.columns if c.key]


def guess(spec: Area, headers: list[str]) -> dict[str, int | None]:
    """Which column each field most likely is, by its header's name."""
    taken: set[int] = set()
    out: dict[str, int | None] = {}
    for c in fields_of(spec):
        names = {normal(a) for a in (c.label, c.key, *c.aliases)}
        hit = next((i for i, h in enumerate(headers) if i not in taken and normal(h) in names), None)
        if hit is not None:
            taken.add(hit)
        out[c.key] = hit
    return out


def can_import(access, spec: Area) -> None:
    if spec.input is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "That list can't be imported.")
    access.require(f"{spec.area}.create")
    access.require(f"{spec.area}.edit")


@router.post("/imports/{kind}/read")
async def read_file(kind: ImportKind, access: CurrentAccess, file: Annotated[UploadFile, File()]) -> dict:
    spec = AREAS[kind]
    can_import(access, spec)
    data = await file.read(MAX_BYTES + 1)
    if len(data) > MAX_BYTES:
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_CONTENT, "The file is over 5 MB. Split it, or take out what isn't needed.")
    rows = [r for r in read_rows(file.filename or "", data) if any(c for c in r)]
    if not rows:
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_CONTENT, "The file is empty.")
    headers, body = rows[0], rows[1:]
    if len(body) > MAX_ROWS:
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_CONTENT, f"The file has {len(body)} rows; import at most {MAX_ROWS} at a time.")
    width = max(len(r) for r in rows)
    headers = [h or f"Column {i + 1}" for i, h in enumerate(headers + [""] * (width - len(headers)))]
    return {
        "fileName": file.filename,
        "headers": headers,
        "rows": [r + [""] * (width - len(r)) for r in body],
        "fields": [{"key": c.key, "label": c.label, "required": c.required, "hint": explain(c)} for c in fields_of(spec)],
        "mapping": guess(spec, headers),
    }


# --- importing -------------------------------------------------------------------------------


class ImportIn(Api):
    # field key -> column index (None: not in the file)
    mapping: dict[str, int | None]
    rows: list[list[str]] = Field(max_length=MAX_ROWS)
    dry_run: bool = True


DATE_FORMATS = ("%Y-%m-%d", "%m/%d/%Y", "%d/%m/%Y", "%d %b %Y", "%d %B %Y", "%b %d, %Y", "%B %d, %Y", "%b %d %Y", "%B %d %Y", "%d-%b-%Y", "%Y/%m/%d")


def to_date(text: str) -> str:
    t = text.strip()
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}[ T].*", t):
        t = t[:10]
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(t, fmt).date().isoformat()
        except ValueError:
            continue
    if re.fullmatch(r"\d{5}(\.0+)?", t):  # an Excel serial number
        from openpyxl.utils.datetime import from_excel

        return from_excel(float(t)).date().isoformat()
    raise ValueError(f"{text!r} isn't a date. Write it as 2027-03-31 or 31 Mar 2027.")


def to_number(text: str, what: str) -> float:
    clean = re.sub(r"[^0-9.\-]", "", text)
    try:
        return float(clean)
    except ValueError:
        raise ValueError(f"{text!r} isn't {what}.") from None


def to_choice(c: Col, text: str) -> str:
    want = normal(text)
    for code, label in (c.choices or {}).items():
        if want in (normal(code), normal(label)):
            return code
    # a label's short form, e.g. "BOC" for "Bureau of Customs (BOC)"
    for code, label in (c.choices or {}).items():
        short = re.search(r"\(([^)]+)\)", label)
        if short and want == normal(short.group(1)):
            return code
    options = ", ".join(dict.fromkeys((c.choices or {}).values()))
    raise ValueError(f"{c.label}: {text!r} isn't one of {options}.")


def to_bool(text: str) -> bool:
    t = normal(text)
    if t in ("yes", "y", "true", "1", "on"):
        return True
    if t in ("no", "n", "false", "0", "off"):
        return False
    raise ValueError(f"{text!r}: write Yes or No.")


def convert(c: Col, text: str) -> Any:
    if c.type == "date":
        return to_date(text)
    if c.type == "int":
        value = to_number(text, "a number of days")
        if not value.is_integer() or value < 0:
            raise ValueError(f"{c.label}: write a whole number of days.")
        return int(value)
    if c.type in ("money", "percent"):
        value = to_number(text, "an amount" if c.type == "money" else "a number")
        if value < 0:
            raise ValueError(f"{c.label} can't be below zero.")
        return value
    if c.type == "bool":
        return to_bool(text)
    if c.type == "choice":
        return to_choice(c, text)
    return text.strip()


def payload_of(spec: Area, row: list[str], mapping: dict[str, int | None]) -> BaseModel:
    data: dict[str, Any] = {}
    for c in fields_of(spec):
        at = mapping.get(c.key)
        text = row[at].strip() if at is not None and 0 <= at < len(row) else ""
        if not text:
            if c.required:
                raise ValueError(f"{c.label} is empty.")
            continue
        data[c.key] = convert(c, text)
    try:
        return spec.input(**data)  # type: ignore[misc]
    except ValidationError as e:
        first = e.errors()[0]
        name = next((c.label for c in spec.columns if c.key == (first["loc"][0] if first["loc"] else "")), "A value")
        raise ValueError(f"{name}: {first['msg'].lower()}.") from None


def fill(spec: Area, rec: Any, payload: BaseModel, known: set[str]) -> None:
    """Sets what the file gave. On an existing record, columns not in the file
    stay as they are."""
    for c in fields_of(spec):
        if c.key not in known:
            continue
        value = getattr(payload, c.key)
        if c.type == "date":
            value = parse_iso_datetime(value)
        elif isinstance(value, str):
            value = value.strip() or None
        setattr(rec, c.attr, value)


@router.post("/imports/{kind}")
def run_import(kind: ImportKind, body: ImportIn, access: CurrentAccess, db: DB, settings: SettingsDep) -> dict:
    spec = AREAS[kind]
    can_import(access, spec)
    missing = [c.label for c in fields_of(spec) if c.required and body.mapping.get(c.key) is None]
    if missing:
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_CONTENT, f"Choose the column for {', '.join(missing)}.")
    known = {k for k, v in body.mapping.items() if v is not None}
    added = updated = 0
    problems: list[dict] = []
    seen: dict[str, int] = {}
    for n, row in enumerate(body.rows, start=2):  # row 1 is the headers
        if not any(c.strip() for c in row):
            continue
        try:
            payload = payload_of(spec, row, body.mapping)
        except ValueError as e:
            problems.append({"row": n, "message": str(e)})
            continue
        key = spec.ident(payload)  # type: ignore[misc]
        if key in seen:
            problems.append({"row": n, "message": f"Same {spec.noun} as row {seen[key]}."})
            continue
        seen[key] = n
        existing = db.scalars(select(spec.model).where(spec.match(payload)).limit(2)).all()  # type: ignore[misc]
        if len(existing) > 1:
            problems.append({"row": n, "message": f"More than one {spec.noun} already matches this row. Edit it in the list instead."})
            continue
        if existing:
            rec = existing[0]
            fill(spec, rec, payload, known)
            rec.updated_at = utcnow()
            updated += 1
        else:
            rec = spec.model(id=new_id())
            fill(spec, rec, payload, set(type(payload).model_fields))
            if not body.dry_run:
                db.add(rec)
            added += 1
    if body.dry_run:
        db.rollback()
    else:
        if added or updated:
            audit(db, access.user, f"{spec.kind}.imported", target=f"{added} added, {updated} updated", target_type=spec.kind)
        db.commit()
        sync_notifications(db, settings)
    return {
        "added": added,
        "updated": updated,
        "problems": problems,
        "unmapped": [c.label for c in fields_of(spec) if c.key not in known],
    }
