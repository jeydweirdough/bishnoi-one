"""Quotation templates: designed here or uploaded as Excel, per issuer, with a
logo, an e-signature and control numbers."""

import io
from datetime import datetime

import pytest
from openpyxl import Workbook, load_workbook
from PIL import Image

from tests.test_customers import company, crew  # noqa: F401 (crew is a fixture)
from tests.test_quotes import PDF, catalogue, ids, line, save  # noqa: F401 (fixtures)

LETTER = {"number": "GL-9", "issuedOn": "2026-01-15", "amountCentavos": 500000}


def png(color="black", size=(80, 30)) -> bytes:
    out = io.BytesIO()
    Image.new("RGB", size, color).save(out, "PNG")
    return out.getvalue()


def pcso_form() -> bytes:
    """Like the PCSO quotation sheet: labels on the left, a products table, a total, a signature line."""
    book = Workbook()
    ws = book.active
    ws.title = "Getmeds Quotation Template"
    ws["A4"] = "No. 17 Vatican Bldg., Las Pinas City"
    ws["A7"] = "PRICE QUOTATION"
    ws["A8"] = "Control No. PCSOG-2026-09-00780"
    ws["B12"], ws["C12"], ws["G12"], ws["H12"] = "PATIENT NAME:", "Old Patient", "Date Issued", datetime(2026, 1, 1)
    ws["B14"], ws["C14"], ws["G14"] = "CONTACT NO.:", 9170000000, "Area Code"
    ws["B16"], ws["C16"] = "ADDRESS:", "Old address"
    ws["B18"], ws["B19"], ws["C19"] = "HOSPITAL/CLINIC", "/CENTER:", "Old hospital"
    for col, head in zip("ABCDEFGHIJ", ["No.", "Generic Name", "Brand Name", "Strength", "Preparation", "Quantity", "Price", "Discount", "Price After Discount", "Total"], strict=True):
        ws[f"{col}22"] = head
    ws["A23"], ws["B23"], ws["F23"], ws["G23"], ws["J23"] = 1, "Old drug", 12, 5316.92, "=F23*G23"
    ws.merge_cells("H24:I24")
    ws["H24"], ws["J24"] = "GRAND TOTAL", "=SUM(J23:J23)"
    ws["A28"] = "We are willing to accept guarantee letter from the PCSO"
    ws["A30"], ws["A31"] = "No. 17 Vatican Drive", "Address"
    ws["A33"] = "FHAYE NORELLA"
    ws["A34"] = "Printed Name/Signature (Authorized Representative)"
    book.create_sheet("Tracker")["A1"] = "not the form"
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


@pytest.fixture
def pcso_quote(client, crew, ids):  # noqa: F811
    acme = company(client, crew["ria"])
    lines = [line(ids["AMX-500"], 2), line(ids["CTZ-10"], 1), line(ids["ORS-SACH"], 3)]
    r = save(client, crew["ria"], acme, lines, payer="PCSO", letter=LETTER, letter_file=PDF)
    assert r.status_code == 201, r.text
    return r.json()


def test_admin_designs_a_form_with_logo_and_signature(client, ana, crew, pcso_quote):  # noqa: F811
    r = client.post(
        "/admin/quote-templates",
        json={"name": "PCSO Getmeds", "kind": "form", "payer": "PCSO", "controlPrefix": "PCSOG", "controlNext": 780, "controlDigits": 5,
              "design": {"signatoryName": "Fhaye Norella", "signatoryTitle": "Customer Happiness Coach"}},
        headers=ana,
    )
    assert r.status_code == 201, r.text
    t = r.json()
    assert t["payerName"] == "PCSO" and t["controlExample"] == "PCSOG-YYYY-MM-00780"
    assert client.post(f"/admin/quote-templates/{t['id']}/images/logo", files={"file": ("logo.png", png("red"), "image/png")}, headers=ana).json()["hasLogo"]
    assert client.post(f"/admin/quote-templates/{t['id']}/images/signature", files={"file": ("sig.png", png(), "image/png")}, headers=ana).json()["hasSignature"]
    bad = client.post(f"/admin/quote-templates/{t['id']}/images/logo", files={"file": ("x.gif", b"GIF89a", "image/gif")}, headers=ana)
    assert bad.status_code == 422

    # the rep sees it for their PCSO quote, and prints with it
    assert [x["name"] for x in client.get(f"/quotes/{pcso_quote['id']}/templates", headers=crew["ria"]).json()["items"]] == ["PCSO Getmeds"]
    doc = client.get(f"/quotes/{pcso_quote['id']}/document", headers=crew["ria"]).json()
    assert doc["controlNumber"].startswith("PCSOG-") and doc["controlNumber"].endswith("-00780")
    assert doc["design"]["acceptance"] == "We are willing to accept a guarantee letter from PCSO."
    assert doc["design"]["signatoryName"] == "Fhaye Norella"
    assert len(doc["lines"]) == 3 and doc["lines"][0]["qty"] == 2
    assert client.get(f"/quote-templates/{t['id']}/images/signature", headers=crew["ria"]).headers["content-type"] == "image/png"
    # the same number every time; the next quote gets the next one
    again = client.get(f"/quotes/{pcso_quote['id']}/document", headers=crew["ria"]).json()
    assert again["controlNumber"] == doc["controlNumber"]
    assert client.get(f"/admin/quote-templates/{t['id']}", headers=ana).json()["controlNext"] == 781


def test_a_template_for_one_issuer_isnt_offered_for_another(client, ana, crew, ids):  # noqa: F811
    client.post("/admin/quote-templates", json={"name": "PCSO only", "payer": "PCSO"}, headers=ana)
    acme = company(client, crew["ria"])
    own = save(client, crew["ria"], acme, [line(ids["AMX-500"])]).json()
    assert client.get(f"/quotes/{own['id']}/templates", headers=crew["ria"]).json()["items"] == []
    assert client.get(f"/quotes/{own['id']}/document", headers=crew["ria"]).status_code == 404
    client.post("/admin/quote-templates", json={"name": "Any quote"}, headers=ana)
    assert [x["name"] for x in client.get(f"/quotes/{own['id']}/templates", headers=crew["ria"]).json()["items"]] == ["Any quote"]


def test_an_uploaded_excel_form_is_filled_in(client, ana, crew, pcso_quote):  # noqa: F811
    t = client.post("/admin/quote-templates", json={"name": "PCSO Excel", "payer": "PCSO", "controlPrefix": "PCSOG", "controlNext": 12}, headers=ana).json()
    r = client.post(f"/admin/quote-templates/{t['id']}/excel", files={"file": ("PCSO TRACKER.xlsx", pcso_form(), "application/octet-stream")}, headers=ana)
    assert r.status_code == 200, r.text
    t = r.json()
    assert t["kind"] == "excel" and t["excel"]["sheet"] == "Getmeds Quotation Template"
    assert t["excel"]["sheets"] == ["Getmeds Quotation Template", "Tracker"]
    found = " | ".join(t["excel"]["found"])
    for part in ("Patient or customer name: C12", "Contact number: C14", "Hospital or clinic: C19", "Date issued: H12", "Control number: A8", "Grand total: J24"):
        assert part in found
    assert t["excel"]["missing"] == []
    client.post(f"/admin/quote-templates/{t['id']}/images/signature", files={"file": ("sig.png", png(), "image/png")}, headers=ana)

    r = client.get(f"/quotes/{pcso_quote['id']}/document.xlsx", headers=crew["ria"])
    assert r.status_code == 200, r.text
    ws = load_workbook(io.BytesIO(r.content)).active
    assert load_workbook(io.BytesIO(r.content)).sheetnames == ["Getmeds Quotation Template"]
    assert ws["C12"].value == acme_name(pcso_quote)
    assert ws["A8"].value.startswith("Control No. PCSOG-") and ws["A8"].value.endswith("-0012")
    # three lines: two rows were added above the total, which moved down
    assert [ws[f"A{r}"].value for r in (23, 24, 25)] == [1, 2, 3]
    assert ws["F23"].value == 2 and ws["G23"].value > 0
    assert ws["H26"].value == "GRAND TOTAL" and ws["J26"].value == pcso_quote["totalCentavos"] / 100
    assert "H26:I26" in {str(m) for m in ws.merged_cells.ranges}
    assert ws["A35"].value == "FHAYE NORELLA"  # the signature line moved down with the rest
    assert ws["A33"].value == "Address"  # the caption above it is left alone
    assert len(ws._images) == 1


def acme_name(quote: dict) -> str:
    return quote["customerName"]


def test_only_admins_set_templates_up(client, crew):  # noqa: F811
    assert client.post("/admin/quote-templates", json={"name": "Mine"}, headers=crew["ria"]).status_code == 403
    assert client.get("/admin/quote-templates", headers=crew["ria"]).status_code == 403


def test_a_control_number_taken_by_another_pc_moves_to_the_next(app, client, ana, crew, ids, pcso_quote):  # noqa: F811
    """Two PCs printing at once: if the next number was just taken, this one takes the one after."""
    from app.models import QuoteControl, QuoteTemplate

    t = client.post("/admin/quote-templates", json={"name": "PCSO", "payer": "PCSO", "controlPrefix": "PCSOG", "controlNext": 780, "controlDigits": 5}, headers=ana).json()
    taken = f"PCSOG-{pcso_quote['createdAt'][:7]}-00780"
    with app.state.sessions() as db:
        other = save(client, crew["ria"], {"id": pcso_quote["customerId"]}, [line(ids["AMX-500"])], payer="PCSO", letter=LETTER, letter_file=PDF).json()["id"]
        db.add(QuoteControl(template_id=t["id"], quote_id=other, number=taken))  # another PC's, a moment ago
        db.commit()
        assert db.get(QuoteTemplate, t["id"]).control_next == 780  # that PC hadn't moved the counter yet
    got = client.get(f"/quotes/{pcso_quote['id']}/document", headers=crew["ria"]).json()["controlNumber"]
    assert got.endswith("-00781")
